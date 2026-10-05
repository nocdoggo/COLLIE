"""Annotation sheets as Excel workbooks, and back.

The samplers (``audit_sample.py``, ``audit_outcomes.py``) write an annotator's sheet as a CSV
file: lines that begin with ``#``, then one table. A spreadsheet program that opens such a file
directly may convert what it shows (``2021-03`` becomes a date, an NDC loses its zeros) and may
save it in another encoding or with another separator. This module gives the annotator the same
sheet as a workbook in which every cell is text, so nothing has to be imported, and turns the
filled workbook back into the CSV layout that the validators read.

* ``to-xlsx`` writes ``<sheet>.xlsx`` from a blank CSV sheet: one worksheet row per CSV row and
  one cell per CSV cell, every cell stored as text and, but for a very long shown one, in the
  text format, so that what is typed stays what was typed. The shown columns are shaded, the
  columns to fill are not, the header row is frozen, a long run of shown columns can be folded
  away, and the columns with a closed set of values get a list to pick from. The list only
  offers its values: no entry is refused and none raises a message, so a list can never block
  a label.
* ``to-csv`` reads a filled workbook and writes the CSV sheet in the layout of the blank: the
  same quoting, the same line ends, ``#`` lines first, then the table. The validators of the
  samplers remain the check of the content; this module changes the file format only, never
  an entry, and prints whatever it wrote otherwise than it was stored.
* A workbook made from a blank and turned back gives the bytes of the blank.

What ``to-csv`` does with a workbook as a spreadsheet program saves it:

* It stops, and writes nothing, when the table's header row or its item ids or their order
  are not the blank's, when a ``#`` line of the blank is missing, when a cell holds a formula
  or an error value (a formula is not an entry, and one that was never calculated has no
  value at all), when cells in the columns to fill are merged (the cells under a merge are
  empty), when text stands to the right of the last column or above the table outside a ``#``
  line, and when more than one worksheet holds the table (``--sheet`` names the one to read).
* It writes as text, and names, every cell stored as a number, a truth value (``TRUE``,
  ``FALSE``), a date or a time. A date becomes ISO text; a date that the workbook shows
  without a day (``2021-06`` typed into a cell that is not text is stored as the first of June
  and shown as ``Jun-21``) becomes the month, ``2021-06``, which is what was entered.
* It names every shown cell and every ``#`` line that differs from the blank, every merged
  range, every cell comment, every other worksheet that holds text, and the empty rows it
  left out.
* A ``#`` line is one line of text. One that is spread over several cells is joined, a line
  break in it becomes a space, and a full-width colon or semicolon becomes the plain one;
  each such line is named. A carriage return inside a cell becomes a line break.
* It does not replace an existing ``--out`` file that holds something else, so that a second
  return cannot take the place of the first unnoticed (``--replace`` allows it).

The module uses ``openpyxl`` and nothing of the study's code, so it runs outside the study's
environment, which does not hold that library::

    uv run --no-project --with openpyxl python analysis/coling/sheet_xlsx.py to-xlsx BLANK.csv OUT.xlsx
    uv run --no-project --with openpyxl python analysis/coling/sheet_xlsx.py to-csv FILLED.xlsx --blank BLANK.csv --out FILLED.csv
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import re
import sys
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from itertools import zip_longest
from pathlib import Path
from typing import Any

FIRST_ENTERED = ("statement_type", "verdict_B", "ok_type", "verdict", "decision")
"""The first column that is filled: on a literal-reading sheet, on an outcome sheet, on a
minimal-pair sheet (task C), on a reference-reading sheet (task E), and on an adjudication
sheet."""
OUTCOME_LISTS = {
    "verdict_B": ("ok", "error", "cannot_tell"),
    "verdict_A": ("ok", "error", "cannot_tell"),
    "true_B_type": ("recovered", "discontinued", "censored", "not_at_risk"),
    "true_A_type": ("recovered", "discontinued", "censored", "not_at_risk"),
}
"""Closed value sets of the outcome sheet (``VERDICTS`` and ``TYPES`` of ``audit_outcomes.py``).
The literal-reading sheet names its own in its ``# values of ...`` lines."""
PAIR_LISTS = {
    name: ("1", "0")
    for name in (
        "ok_type",
        "ok_interval",
        "ok_certainty",
        "ok_stale",
        "ok_distractors",
        "minimal",
        "attested",
        "natural",
    )
}
"""Closed value sets of the minimal-pair sheet: its eight marks (``ERROR_CHECKS`` and
``natural`` of ``minimal_pairs.py``), which its ``# values of ...`` lines describe in words."""
SINGLE_VALUED = (
    "statement_type",
    "abstain",
    "abstain_reason",
    "certainty",
    "hard",
    "verdict",
    "start_differs",
)
"""Columns that hold one value of a closed set which the sheet names in a ``# values of ...``
line: those of the literal-reading sheet, and the verdict and the start flag of the
reference-reading sheet."""
WIDE = 45
NARROW = 12
ENTERED = 18
COMMENT_HEIGHT = 13.5
"""Column widths in characters and the height of a ``#`` line in points."""
FOLD_ABOVE = 12
"""A sheet with more shown columns than this gets them in one group that can be folded."""
TEXT_LIMIT = 32767
"""The longest text a worksheet cell holds; ``openpyxl`` cuts a longer one without a word."""
TEXT_SHOWN = 255
"""The longest text that a cell of the text format is sure to show: older versions of Excel
show a longer one as a row of ``#``. A longer shown cell keeps the general format; it is
stored as text all the same, and no cell to fill is that long."""
ESCAPED = re.compile(r"_x([0-9A-Fa-f]{4})_")
"""How a workbook writes a character that XML cannot hold: a carriage return is ``_x000D_``."""
ESCAPED_UNDERSCORE = b"_x005F_"
"""What a workbook writes before literal text of that form. ``openpyxl`` drops it, so that the
literal text and the escape cannot be told apart once the workbook is loaded."""
PLAIN_FORMS = str.maketrans({0xFF1A: ":", 0xFF1B: ";"})
"""In a ``#`` line, the full-width colon and semicolon of an input method as the plain ones."""
SHOWN_LENGTH = 60
"""The longest text of a cell that a message quotes."""


@dataclass(frozen=True)
class Sheet:
    """A CSV sheet as rows of cells, with the layout needed to write it back unchanged."""

    rows: list[list[str]]
    header_at: int
    """Index of the table's header row; the rows before it are the ``#`` lines."""
    quoted_comments: bool
    """Whether a ``#`` line is written as a quoted CSV cell (else as a bare line)."""
    newline: str

    @property
    def header(self) -> list[str]:
        return self.rows[self.header_at]


@dataclass
class Page:
    """One worksheet read as rows of text, with what was found while reading it."""

    title: str
    rows: list[list[str]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    """Cells written otherwise than they were stored."""
    problems: list[str] = field(default_factory=list)
    """Cells that hold no entry: a formula, an error value."""
    merged: list[tuple[int, int, int, int]] = field(default_factory=list)
    """Merged ranges: first row, first column, last row, last column, counted from 1."""


def read_csv_sheet(path: Path) -> Sheet:
    """The rows of a CSV sheet and how its ``#`` lines and line ends are written."""
    text = path.read_bytes().decode("utf-8")
    if text.startswith(chr(0xFEFF)):
        raise ValueError(f"{path}: starts with a byte-order mark; a sampler did not write it")
    first_line = text.split("\n", 1)[0]
    newline = "\r\n" if first_line.endswith("\r") else "\n"
    quoted = text.startswith('"#')
    lines = text.split(newline)
    mark = '"#' if quoted else "#"
    header_at = next((n for n, line in enumerate(lines) if not line.startswith(mark)), len(lines))
    if quoted:
        comments = list(csv.reader(io.StringIO(newline.join(lines[:header_at]), newline="")))
    else:
        comments = [[line] for line in lines[:header_at]]
    if any(len(row) != 1 for row in comments):
        raise ValueError(f"{path}: a # line is not one cell")
    table = list(csv.reader(io.StringIO(newline.join(lines[header_at:]), newline="")))
    if not table or not table[0]:
        raise ValueError(f"{path}: no table under the # lines")
    return Sheet(comments + table, header_at, quoted, newline)


def write_csv_sheet(sheet: Sheet) -> str:
    """The text of a CSV sheet: ``#`` lines as the blank writes them, every table cell quoted."""
    out = io.StringIO(newline="")
    quoting = csv.writer(out, quoting=csv.QUOTE_ALL, lineterminator=sheet.newline)
    for n, row in enumerate(sheet.rows):
        if n < sheet.header_at and not sheet.quoted_comments:
            out.write((row[0] if row else "") + sheet.newline)
        else:
            quoting.writerow(row)
    return out.getvalue()


def value_lists(sheet: Sheet) -> dict[str, tuple[str, ...]]:
    """The closed value set of each column that has one."""
    found: dict[str, tuple[str, ...]] = {}
    for row in sheet.rows[: sheet.header_at]:
        match = re.match(r"#\s*values of (\w+):\s*(.+)$", row[0] if row else "")
        if match and match.group(1) in SINGLE_VALUED:
            text = match.group(2).split("(")[0]
            parts = [p.strip() for p in re.split(r"\||\bor\b", text) if p.strip()]
            found[match.group(1)] = tuple(parts)
    found |= {k: v for k, v in (OUTCOME_LISTS | PAIR_LISTS).items() if k in sheet.header}
    return {k: v for k, v in found.items() if k in sheet.header}


def first_entered(header: list[str]) -> int:
    """The index of the first column an annotator fills."""
    for name in FIRST_ENTERED:
        if name in header:
            return header.index(name)
    return len(header)


def to_xlsx(blank: Path, out: Path) -> None:
    """Write a workbook with the cells of a blank CSV sheet, every cell stored as text."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    sheet = read_csv_sheet(blank)
    if any(len(cell) > TEXT_LIMIT for row in sheet.rows for cell in row):
        raise ValueError(f"{blank}: a cell is longer than a worksheet cell ({TEXT_LIMIT})")
    header, entered = sheet.header, first_entered(sheet.header)
    book = Workbook()
    book.properties.creator = ""
    book.properties.lastModifiedBy = ""
    page = book.active
    page.title = "sheet"
    shown_fill = PatternFill("solid", start_color="FFE7E6E6")
    head_fill = PatternFill("solid", start_color="FFBDD7EE")
    fill_fill = PatternFill("solid", start_color="FFFFF2CC")
    top = Alignment(vertical="top", wrap_text=True)
    for r, row in enumerate(sheet.rows, start=1):
        width = len(header) if r > sheet.header_at else 1
        for c in range(1, width + 1):
            cell = page.cell(row=r, column=c)
            text = row[c - 1] if c <= len(row) else ""
            cell.number_format = "@" if len(text) <= TEXT_SHOWN else "General"
            if text:
                cell.value = text
                cell.data_type = "s"
            if r <= sheet.header_at:
                cell.font = Font(italic=True, color="FF595959")
                cell.alignment = Alignment(vertical="top", wrap_text=False)
            elif r == sheet.header_at + 1:
                cell.font = Font(bold=True)
                cell.fill = head_fill if c <= entered else fill_fill
                cell.alignment = top
            else:
                if c <= entered:
                    cell.fill = shown_fill
                cell.alignment = top
    first, last = sheet.header_at + 2, len(sheet.rows)
    for name, values in value_lists(sheet).items():
        letter = get_column_letter(header.index(name) + 1)
        # no message for a value outside the list: a list offers its values and refuses nothing
        rule = DataValidation(
            type="list",
            formula1='"' + ",".join(values) + '"',
            allow_blank=True,
            showErrorMessage=False,
            errorStyle="warning",
            errorTitle=name,
            error="This value is not one of: " + ", ".join(values),
        )
        rule.add(f"{letter}{first}:{letter}{last}")
        page.add_data_validation(rule)
    body = sheet.rows[sheet.header_at + 1 :]
    for c, name in enumerate(header, start=1):
        longest = max((len(row[c - 1]) for row in body if c <= len(row)), default=0)
        wanted = max(longest, len(name)) + 2
        if c > entered:
            wanted = max(wanted, ENTERED)
        page.column_dimensions[get_column_letter(c)].width = max(NARROW, min(WIDE, wanted))
    for r in range(1, sheet.header_at + 1):
        page.row_dimensions[r].height = COMMENT_HEIGHT
    if FOLD_ABOVE < entered < len(header):
        # many shown columns: they can be folded away, to bring the columns to fill beside the id
        for c in range(2, entered + 1):
            page.column_dimensions[get_column_letter(c)].outlineLevel = 1
    page.freeze_panes = page.cell(row=sheet.header_at + 2, column=2)
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)


def column_letter(column: int) -> str:
    """The letter of a worksheet column counted from 1 (``28`` is ``AB``)."""
    letters = ""
    while column > 0:
        column, rest = divmod(column - 1, 26)
        letters = chr(ord("A") + rest) + letters
    return letters


def clip(text: str, start: int = 0) -> str:
    """The text of a cell for a message: in quotes, and no more than ``SHOWN_LENGTH`` of its
    characters, from ``start`` on."""
    before = "..." if start else ""
    after = "..." if len(text) > start + SHOWN_LENGTH else ""
    return before + repr(text[start : start + SHOWN_LENGTH]) + after


def difference(mine: str, theirs: str) -> str:
    """Two texts that differ, for a message: both from the start, or from shortly before their
    first difference when that lies further on."""
    pairs = enumerate(zip(mine, theirs, strict=False))
    at = next((n for n, (a, b) in pairs if a != b), min(len(mine), len(theirs)))
    start = 0 if at < SHOWN_LENGTH * 2 // 3 else at - SHOWN_LENGTH // 3
    return f"{clip(mine, start)}, the blank holds {clip(theirs, start)}"


def date_letters(number_format: str) -> str:
    """The letters of a date format that stand for parts of a date: its first section in lower
    case, without quoted text, codes in brackets and escaped characters."""
    code = re.sub(r'"[^"]*"|\[[^\]]*\]|[\\_*].', "", number_format)
    return code.split(";")[0].lower()


def clock_text(value: dt.time) -> str:
    return value.strftime("%H:%M:%S" if value.second else "%H:%M")


def stored_text(value: Any, number_format: str) -> tuple[str, str]:
    """The text of a cell that is not stored as text, and what it is stored as.

    A truth value is written as the workbook shows it, not as ``1`` or ``0``. A date is written
    as ISO text; one that the workbook shows as a month, without a day and without a time, and
    stores as the first of that month, is written as the month (``2021-06``): for an ``end``
    that is the last day of June, and the first of June would be another label.
    """
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), "a truth value"
    if isinstance(value, dt.time):
        return clock_text(value), f"a time (shown as {number_format})"
    if isinstance(value, dt.datetime | dt.date):
        day = value.date() if isinstance(value, dt.datetime) else value
        clock = value.time().replace(microsecond=0) if isinstance(value, dt.datetime) else dt.time()
        letters = set(date_letters(number_format))
        if clock != dt.time():
            text = f"{day.isoformat()} {clock_text(clock)}"
        elif "m" in letters and not letters & set("dhs") and day.day == 1:
            text = day.isoformat()[:7]
        else:
            text = day.isoformat()
        return text, f"a date (shown as {number_format})"
    if isinstance(value, dt.timedelta):
        return str(value), f"a duration (shown as {number_format})"
    whole = isinstance(value, float) and value.is_integer()
    return str(int(value) if whole else value), f"a number (shown as {number_format})"


def cell_text(cell: Any, page: Page, unsure: bool) -> str:
    """A worksheet cell as text; what is not stored as text goes to the page's notes, and what
    is no entry at all to its problems."""
    ref = cell.coordinate
    if cell.comment is not None:
        text = clip(cell.comment.text or "")
        page.notes.append(f"{ref}: carries a comment, which is not written: {text}")
    value = cell.value
    if value is None:
        return ""
    if cell.data_type == "f":
        formula = clip(str(getattr(value, "text", value)))
        page.problems.append(f"{ref}: holds a formula, which is not an entry: {formula}")
        return ""
    if cell.data_type == "e":
        page.problems.append(f"{ref}: holds an error value, which is not an entry: {value}")
        return ""
    if not isinstance(value, str):
        text, stored = stored_text(value, cell.number_format)
        page.notes.append(f"{ref}: stored as {stored}, written as the text {clip(text)}")
        return text
    text = value
    if ESCAPED.search(text):
        if unsure:
            page.problems.append(
                f"{ref}: holds text of the form _x000D_, which cannot be read with certainty: "
                + clip(text)
            )
            return ""
        text = ESCAPED.sub(lambda match: chr(int(match.group(1), 16)), text)
        page.notes.append(f"{ref}: escaped characters are written as the characters they name")
    if "\r" in text:
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        page.notes.append(f"{ref}: a carriage return inside the cell is written as a line break")
    return text


def read_xlsx(path: Path) -> list[Page]:
    """Every worksheet of a workbook as rows of text, with what was found while reading it.

    The workbook is loaded with its formulas, not with their last results: a formula that was
    never calculated has no result, and would read as an empty cell.
    """
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        with zipfile.ZipFile(path) as archive:
            parts = [name for name in archive.namelist() if name.endswith(".xml")]
            unsure = any(ESCAPED_UNDERSCORE in archive.read(name) for name in parts)
        book = load_workbook(path)
    except (zipfile.BadZipFile, InvalidFileException, KeyError) as problem:
        raise ValueError(f"{path}: not an Excel workbook (.xlsx): {problem}") from None
    pages: list[Page] = []
    for sheet in book.worksheets:
        page = Page(sheet.title)
        for cells in sheet.iter_rows():
            row = [cell_text(cell, page, unsure) for cell in cells]
            while row and row[-1] == "":
                row.pop()
            page.rows.append(row)
        while page.rows and not page.rows[-1]:
            page.rows.pop()
        for merged in sheet.merged_cells.ranges:
            first_column, first_row, last_column, last_row = merged.bounds
            page.merged.append((first_row, first_column, last_row, last_column))
        pages.append(page)
    return pages


def comment_key(line: str) -> str:
    """The key of a ``# key: value`` line, as the validators read it."""
    return line.strip().lstrip("# ").partition(":")[0].strip()


def comment_value(line: str) -> str:
    return line.strip().lstrip("# ").partition(":")[2].strip()


def comment_line(row: list[str]) -> str:
    """A ``#`` line as one line of text: the cells of its row joined by a space, a line break
    as a space, a full-width colon or semicolon as the plain one. A line that needs none of
    this is the first cell as it stands."""
    line = row[0].lstrip()
    others = [cell.strip() for cell in row[1:] if cell.strip()]
    if others:
        line = " ".join([line.rstrip(), *others])
    if "\n" in line:
        line = " ".join(part.strip() for part in line.split("\n") if part.strip())
    return line.translate(PLAIN_FORMS)


def comment_lines(
    rows: list[list[str]], blank: list[str], notes: list[str], problems: list[str]
) -> list[str]:
    """The ``#`` lines of a filled sheet, from the worksheet rows above the table.

    A line of the blank that names a value (the task, the guide, the value sets) is expected
    as it was; one that names none (the times) is the annotator's to fill, and may be repeated.
    """
    lines: list[str] = []
    for n, row in enumerate(rows, start=1):
        if not any(cell.strip() for cell in row):
            continue
        if not row[0].lstrip().startswith("#"):
            text = clip(" ".join(cell for cell in row if cell.strip()))
            problems.append(f"row {n}, above the table, is not a # line: {text}")
            continue
        line = comment_line(row)
        if line != row[0] or len(row) > 1:
            notes.append(f"row {n}: the # line is written as {line!r}")
        lines.append(line)
        same_key = [b for b in blank if comment_key(b) == comment_key(line)]
        if not same_key:
            notes.append(f"row {n}: a # line that the blank does not hold: {clip(line)}")
        elif line.strip() not in [b.strip() for b in same_key] and comment_value(same_key[0]):
            notes.append(f"row {n}: # line changed: {line!r}, the blank holds {same_key[0]!r}")
    seen = {comment_key(line) for line in lines}
    missing = sorted({comment_key(b) for b in blank} - seen)
    if missing:
        problems.append("# lines missing: " + "; ".join(f"# {key}" for key in missing))
    return lines


def header_problem(rows: list[list[str]], header: list[str]) -> str:
    """Why no row of the worksheet is the blank's header row."""
    found = next((row for row in rows if row and not row[0].lstrip().startswith("#")), None)
    if found is None:
        return "no table under the # lines"
    pairs = enumerate(zip_longest(found, header, fillvalue=""), start=1)
    changed = [
        f"{column_letter(c)} holds {got!r} for {want!r}" for c, (got, want) in pairs if got != want
    ]
    more = f"; and {len(changed) - 3} more" if len(changed) > 3 else ""
    return "the header row of the table is not the blank's: " + "; ".join(changed[:3]) + more


def id_problem(table: list[tuple[int, list[str]]], theirs: list[str]) -> str:
    """Why the item ids of a filled sheet, in their order, are not the blank's."""
    mine = [row[0] for _, row in table]
    parts = [f"{i!r} is missing" for i in theirs if i not in mine]
    parts += [f"row {n} holds {row[0]!r}, not an item" for n, row in table if row[0] not in theirs]
    parts += [
        f"{i!r} is there {k} times" for i, k in Counter(mine).items() if k > 1 and i in theirs
    ]
    if not parts:
        at = next(k for k, i in enumerate(theirs) if mine[k] != i)
        here = f"row {table[at][0]} holds {mine[at]!r}, the blank {theirs[at]!r}"
        parts = [f"the rows are in another order ({here})"]
    more = f"; and {len(parts) - 5} more" if len(parts) > 5 else ""
    return "the item ids, or their order, are not the blank's: " + "; ".join(parts[:5]) + more


def to_csv(filled: Path, blank: Path, sheet: str | None = None) -> tuple[str, list[str], list[str]]:
    """The CSV text of a filled workbook in the blank's layout, with problems (which stop the
    conversion: the text is then empty) and notes (which do not). ``sheet`` names the worksheet
    to read; without it the one that holds the blank's header row is read."""
    base = read_csv_sheet(blank)
    width = len(base.header)
    pages = read_xlsx(filled)
    titles = ", ".join(repr(page.title) for page in pages)
    if sheet is not None:
        chosen = [page for page in pages if page.title == sheet]
        if not chosen:
            return "", [f"no worksheet named {sheet!r}; the workbook holds {titles}"], []
    else:
        chosen = [p for p in pages if any(row[:width] == base.header for row in p.rows)]
        if len(chosen) > 1:
            held = ", ".join(repr(page.title) for page in chosen)
            return (
                "",
                [f"more than one worksheet holds the table ({held}); name one with --sheet"],
                [],
            )
        chosen = chosen or pages[:1]
    if not chosen:
        return "", ["the workbook holds no worksheet"], []
    page = chosen[0]
    problems, notes = list(page.problems), list(page.notes)
    if page is not pages[0]:
        notes.append(f"the sheet is read from worksheet {page.title!r}, which is not the first")
    for other in pages:
        if other is not page and other.rows:
            cells = sum(1 for row in other.rows for cell in row if cell.strip())
            notes.append(f"worksheet {other.title!r} is not read ({cells} cells with text)")
    rows = page.rows
    header_at = next((n for n, row in enumerate(rows) if row[:width] == base.header), None)
    if header_at is None:
        return "", [*problems, header_problem(rows, base.header)], notes
    blank_lines = [row[0] if row else "" for row in base.rows[: base.header_at]]
    comments = comment_lines(rows[:header_at], blank_lines, notes, problems)
    table: list[tuple[int, list[str]]] = []
    empty = 0
    for n, row in enumerate(rows[header_at:], start=header_at + 1):
        if not any(cell.strip() for cell in row):
            empty += 1
            continue
        for c, cell in enumerate(row[width:], start=width + 1):
            ref = f"{column_letter(c)}{n}"
            if cell.strip():
                problems.append(f"{ref}: text to the right of the last column: {clip(cell)}")
            elif cell:
                notes.append(f"{ref}: a cell of spaces to the right of the last column is left out")
        table.append((n, (row + [""] * width)[:width]))
    if empty:
        notes.append(f"{empty} empty rows between the rows of the table are left out")
    blank_table = [(row + [""] * width)[:width] for row in base.rows[base.header_at :]]
    if [row[0] for _, row in table] != [row[0] for row in blank_table]:
        problems.append(id_problem(table[1:], [row[0] for row in blank_table[1:]]))
        return "", problems, notes
    entered = first_entered(base.header)
    for (n, mine), theirs in zip(table[1:], blank_table[1:], strict=True):
        for c in range(entered):
            if mine[c] != theirs[c]:
                notes.append(
                    f"shown cell changed: item {theirs[0]}, column {base.header[c]} "
                    f"({column_letter(c + 1)}{n}): {difference(mine[c], theirs[c])}"
                )
    for first_row, first_column, last_row, last_column in page.merged:
        span = f"{column_letter(first_column)}{first_row}:{column_letter(last_column)}{last_row}"
        if last_row > header_at + 1 and entered < last_column and first_column <= width:
            problems.append(
                f"merged cells {span} in the columns to fill: the cells under a merge are "
                "empty, so each of them has to be filled on its own"
            )
        else:
            notes.append(f"merged cells {span}: only the first cell of a merge holds text")
    lines = [[line] for line in comments] + [row for _, row in table]
    out = Sheet(lines, len(comments), base.quoted_comments, base.newline)
    return ("" if problems else write_csv_sheet(out)), problems, notes


def same_file(one: Path, other: Path) -> bool:
    if one.resolve() == other.resolve():
        return True
    return one.exists() and other.exists() and one.samefile(other)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=(__doc__ or "").split("\n\n")[0])
    sub = ap.add_subparsers(dest="command", required=True)
    one = sub.add_parser("to-xlsx", help="a workbook from a blank CSV sheet")
    one.add_argument("blank", type=Path)
    one.add_argument("out", type=Path)
    two = sub.add_parser("to-csv", help="the CSV sheet of a filled workbook")
    two.add_argument("filled", type=Path)
    two.add_argument("--blank", type=Path, required=True)
    two.add_argument("--out", type=Path, required=True)
    two.add_argument("--sheet", help="the worksheet to read, when several hold the table")
    for command in (one, two):
        command.add_argument("--replace", action="store_true", help="write over an existing file")
    args = ap.parse_args(argv)
    if same_file(args.out, args.blank):
        print("error: --out is the blank sheet", file=sys.stderr)
        return 2
    if args.command == "to-xlsx":
        if args.out.exists() and not args.replace:
            # it may be a filled workbook: a blank one never takes its place unasked
            print(
                f"error: {args.out} exists; it is not replaced without --replace", file=sys.stderr
            )
            return 2
        try:
            to_xlsx(args.blank, args.out)
            back, problems, _ = to_csv(args.out, args.blank)
        except ValueError as problem:
            print(f"error: {problem}", file=sys.stderr)
            return 1
        if problems or back.encode("utf-8") != args.blank.read_bytes():
            args.out.unlink()
            print("error: the workbook does not give back the blank sheet", file=sys.stderr)
            return 1
        print(f"wrote {args.out}")
        return 0
    if same_file(args.out, args.filled):
        print("error: --out is the filled workbook", file=sys.stderr)
        return 2
    try:
        text, problems, notes = to_csv(args.filled, args.blank, args.sheet)
    except (OSError, ValueError) as problem:
        print(f"error: {problem}", file=sys.stderr)
        return 1
    for note in notes:
        print(f"note: {note}")
    for problem in problems:
        print(f"error: {problem}", file=sys.stderr)
    if problems:
        print(f"error: {args.out} is not written ({len(problems)} problems)", file=sys.stderr)
        return 1
    data = text.encode("utf-8")
    if args.out.is_dir():
        print(f"error: --out is a directory: {args.out}", file=sys.stderr)
        return 2
    if args.out.exists() and args.out.read_bytes() == data:
        print(f"unchanged: {args.out} already holds this sheet ({len(notes)} notes)")
        return 0
    if args.out.exists() and not args.replace:
        # the file may be an earlier return: a second one never takes its place unnoticed
        print(
            f"error: {args.out} exists and holds something else; it is not replaced "
            "(give another --out, or --replace)",
            file=sys.stderr,
        )
        return 2
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_bytes(data)
    print(f"wrote {args.out} ({len(notes)} notes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
