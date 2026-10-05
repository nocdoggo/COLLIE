"""Tests for the workbook form of the annotation sheets (``sheet_xlsx.py``).

Blank sheets in the two layouts of the samplers are written to a temporary directory: the
literal-reading sheet (quoted ``#`` lines, value sets named in them, two sitting lines) and the
outcome sheet (bare ``#`` lines, two session lines that end in a space). Covered:

* the round trip: a workbook from a blank and back gives the bytes of the blank, for the two
  layouts, for a table without ``#`` lines, for cells that a spreadsheet program would convert
  or that CSV has to quote, and for the study's own blank sheets when they are on disk;
* the workbook that ``to-xlsx`` writes: every cell text, the lists and what they allow, the
  frozen rows and the folded columns, no name in its properties;
* workbooks as a spreadsheet program stores them, written here as XML by hand: shared and inline
  strings, runs of formatted text, numbers, truth values, dates, months, times, formulas with
  and without a result, error values, escaped characters, hidden rows, merged cells, further
  worksheets, and what an annotator may do to the rows, the columns and the ``#`` lines;
* the stops and the notes of ``to-csv``, and what the command does with an existing file;
* the two validators of the study on converted sheets, cell by cell (they run in the study's
  environment, as a separate process; skipped when that interpreter is not found: it is the
  one ``COLING_STUDY_PYTHON`` names, or the ``.venv`` of this checkout or of the checkout it
  is a worktree of);
* the workbooks that were handed out, when they are on disk (read only).

No file of the study is written, and nothing sealed is read.

Run, outside the study's environment (which does not hold ``openpyxl``; there the file is
skipped)::

    uv run --no-project --with openpyxl --with pytest python -m pytest analysis/coling/test_sheet_xlsx.py -p no:cacheprovider -q
"""

from __future__ import annotations

import ast
import csv
import datetime as dt
import io
import json
import os
import re
import subprocess
import zipfile
from collections.abc import Callable
from pathlib import Path
from xml.etree import ElementTree
from xml.sax.saxutils import escape

import pytest

pytest.importorskip("openpyxl")

from openpyxl import load_workbook
from openpyxl.comments import Comment

from analysis.coling import sheet_xlsx as X

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
AUDIT = HERE / "out" / "audit"
OUTCOMES = HERE / "out" / "audit_outcomes"
HANDOUT = ROOT / "external_data" / "annotation" / "handout_2026-10-02"
REAL_BLANKS = (
    AUDIT / "pilot_A1.csv",
    AUDIT / "pilot_A2.csv",
    AUDIT / "check_A1.csv",
    AUDIT / "check_A2.csv",
    OUTCOMES / "outcome_A1_train.csv",
    OUTCOMES / "outcome_A2_train.csv",
    OUTCOMES / "outcome_train_traces.csv",
)
HANDED_OUT = (
    ("A1/pilot_A1.xlsx", AUDIT / "pilot_A1.csv"),
    ("A2/pilot_A2.xlsx", AUDIT / "pilot_A2.csv"),
    ("A1/outcome_A1_train.xlsx", OUTCOMES / "outcome_A1_train.csv"),
    ("A2/outcome_A2_train.xlsx", OUTCOMES / "outcome_A2_train.csv"),
    ("A1/outcome_train_traces.xlsx", OUTCOMES / "outcome_train_traces.csv"),
    ("A2/outcome_train_traces.xlsx", OUTCOMES / "outcome_train_traces.csv"),
)

LITERAL_SHOWN = (
    "item_id",
    "drug",
    "company",
    "presentation",
    "therapeutic_category",
    "initial_posting_date",
    "type_of_update",
    "date_of_update",
    "availability_information",
    "related_information",
    "reason_for_shortage",
)
LITERAL_ENTERED = (
    "statement_type",
    "start",
    "end",
    "abstain",
    "abstain_reason",
    "certainty",
    "quote",
    "distractor_roles",
    "distractor_quotes",
    "hard",
    "note",
)
OUTCOME_SHOWN = (
    "audit_id",
    "generic",
    "company",
    "presentation",
    "statement_text",
    "statement_date",
    "c0",
    "followed_to",
    "class_strip",
    *(
        f"derived_{d}_{part}"
        for d in "BA"
        for part in ("type", "lower", "upper", "width", "censor", "exit")
    ),
    "group_n_presentations",
    *(f"group_{d}_{part}" for d in "BA" for part in ("type", "lower", "upper")),
)
OUTCOME_ENTERED = (
    "verdict_B",
    "verdict_A",
    "codes_B",
    "codes_A",
    *(f"true_{d}_{part}" for d in "BA" for part in ("type", "lower", "upper")),
    "note_codes",
    "dou_first_recovered",
    "note",
)
LITERAL_META = (
    ("task", "A literal reading, pilot (AUDIT_GUIDE.md sections 3 and 4)"),
    ("sheet", "pilot"),
    ("annotator", "A1"),
    ("seed", "20261001"),
    ("guide", "v1 draft sha256 0000000000000000"),
    ("items", "4"),
    ("values of statement_type", "recovery | next_delivery | depletion | discontinuation | none"),
    (
        "values of start and end",
        "YYYY-MM-DD, or YYYY-MM for a whole month; both blank when abstaining",
    ),
    ("values of abstain", "1 or 0"),
    ("values of abstain_reason", "tbd | no_date | vague | no_statement (only when abstain is 1)"),
    ("values of certainty", "asserted | estimated | undetermined | no_statement"),
    (
        "values of distractor_roles",
        "expiry; depletion; discontinuation; onset; past; other (one quote per role)",
    ),
    ("values of hard", "1 or 0 (say why in note)"),
    ("sitting_start", ""),
    ("sitting_end", ""),
)
OUTCOME_META = (
    ("task", "B outcome audit (AUDIT_GUIDE.md section 6)"),
    ("auditor", "A1"),
    ("half", "train"),
    ("guide", "v1 draft sha256 0000000000000000"),
    ("followed_to", "2022-10-06"),
    ("items", "3"),
    ("session_start", ""),
    ("session_end", ""),
)
LONG = "a long cell, " * 2000
"""26,000 characters: under the limit of a worksheet cell, far over what a screen shows."""
LITERAL_ITEMS = (
    (
        "I000000001",
        'Propofol, "Injectable" Emulsion; 1%',
        "Hikma (formerly West-Ward)",
        "10 mg/mL – 20 mL vial (NDC 0641-6194-10)",
        "Anesthesia;Neurology",
        "2020-04-10",
        "New",
        "2020-07-23",
        "On backorder. Next release August 2021.",
        "Check wholesalers for inventory.",
        "Demand increase for the drug",
    ),
    (
        "I000000002",
        "Sulfasalazine Tablets",
        "  a company with spaces around  ",
        "bottle of 300\nsecond line of the cell",
        "Gastroenterology",
        "(blank)",
        "Revised",
        "2021-03-12",
        "Next Delivery: March 2021; Estimated recovery: Q4 2021",
        "the manufacturer’s estimate — 供給 été",
        "#1 reason: a cell that begins with the sign of a # line",
    ),
    (
        "I000000003",
        "'=cmd|' /C calc'!A0",
        "'+1 lot",
        "'-- none --",
        "'@wholesaler",
        "0012",
        "TRUE",
        "2021-06-30",
        "2021-03",
        "1e5",
        "03/04/2021",
    ),
    (
        "I000000004",
        "Furosemide Injection",
        "=SUM(1,2)",
        "20 mg per 2 mL, single dose vial (NDC 63323-280-02)",
        "Cardiovascular",
        "2020-04-07",
        "Revised",
        "2021-07-22",
        "Supply TBD; lots exp 11/2021, made 2020",
        LONG,
        "09:30",
    ),
)
OUTCOME_ITEMS = (
    ("Q0000001", "Latanoprost Ophthalmic Solution 0.005%", "a company", "2.5 mL fill – 1 bottle"),
    ("Q0000002", "Ketoprofen Capsules", "the maker’s name, Inc.", "'-75 mg, 100 count"),
    ("Q0000003", "Timolol Maleate", "another company", "15 mL bottle (NDC 61314-227-15)"),
)


# --------------------------------------------------------------------------------------------
# Blank sheets as the samplers write them
# --------------------------------------------------------------------------------------------


def literal_text(items: tuple[tuple[str, ...], ...] = LITERAL_ITEMS, newline: str = "\n") -> str:
    """A literal-reading sheet: quoted ``#`` lines, then the table, every cell quoted."""
    out = io.StringIO(newline="")
    writer = csv.writer(out, lineterminator=newline, quoting=csv.QUOTE_ALL)
    for key, value in LITERAL_META:
        writer.writerow([f"# {key}: {value}".rstrip()])
    writer.writerow([*LITERAL_SHOWN, *LITERAL_ENTERED])
    for item in items:
        writer.writerow([*item, *[""] * len(LITERAL_ENTERED)])
    return out.getvalue()


def outcome_text(items: tuple[tuple[str, ...], ...] = OUTCOME_ITEMS) -> str:
    """An outcome sheet: bare ``# key: value`` lines, then the table, every cell quoted."""
    out = io.StringIO(newline="")
    for key, value in OUTCOME_META:
        out.write(f"# {key}: {value}\n")
    writer = csv.writer(out, lineterminator="\n", quoting=csv.QUOTE_ALL)
    writer.writerow([*OUTCOME_SHOWN, *OUTCOME_ENTERED])
    for item in items:
        shown = [*item, "", "2019-10-07", "2019-10-20", "2022-10-06", "O O | A R R", "recovered"]
        shown += ["2019-12-29", "2020-04-02", "95"]
        shown += [""] * (len(OUTCOME_SHOWN) - len(shown))
        writer.writerow([*shown, *[""] * len(OUTCOME_ENTERED)])
    return out.getvalue()


def table_text() -> str:
    """A table without ``#`` lines, as the traces of the outcome audit; one of its rows begins
    with the sign of a ``#`` line."""
    out = io.StringIO(newline="")
    writer = csv.writer(out, lineterminator="\n", quoting=csv.QUOTE_ALL)
    writer.writerow(["audit_id", "capture", "status"])
    writer.writerows([["Q1", "20191229133343", ""], ["#2", "", "U"], ["Q1", "# no line", "#"]])
    return out.getvalue()


def written(directory: Path, name: str, text: str) -> Path:
    path = directory / name
    path.write_bytes(text.encode("utf-8"))
    return path


@pytest.fixture
def literal(tmp_path: Path) -> Path:
    return written(tmp_path, "pilot_A1.csv", literal_text())


@pytest.fixture
def outcome(tmp_path: Path) -> Path:
    return written(tmp_path, "outcome_A1_train.csv", outcome_text())


def parsed(text: str) -> dict[str, dict[str, str]]:
    """The table of a CSV sheet: the cells of each item by column name."""
    rows = list(csv.reader(io.StringIO(text, newline="")))
    head = next(n for n, row in enumerate(rows) if row and not row[0].startswith("#"))
    return {row[0]: dict(zip(rows[head], row, strict=True)) for row in rows[head + 1 :]}


def comments(text: str) -> list[str]:
    """The ``#`` lines of a CSV sheet as its validators read them: ``key: value``."""
    lines = []
    for row in csv.reader(io.StringIO(text, newline="")):
        if not row or not row[0].startswith("#"):
            break
        lines.append(",".join(cell for cell in row if cell).lstrip("# "))
    return lines


# --------------------------------------------------------------------------------------------
# Workbooks as a spreadsheet program stores them, written by hand
# --------------------------------------------------------------------------------------------

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE = "http://schemas.openxmlformats.org/package/2006"
SHEET_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml"
FORMATS = (
    "General",
    "@",
    "mm-dd-yy",
    "yyyy\\-mm\\-dd",
    "mmm-yy",
    "h:mm",
    "d-mmm",
    "0%",
    'yyyy"年"m"月"',
    "m/d/yy h:mm",
    "[h]:mm:ss",
    'yyyy"年"m"月"d"日"',
    "h:mm:ss",
    "[$-409]mmmm\\ yyyy;@",
)
"""The number formats of the cells below; a cell names one of them."""
BUILT_IN = {"General": 0, "@": 49, "mm-dd-yy": 14, "mmm-yy": 17, "h:mm": 20, "d-mmm": 16, "0%": 9}
BUILT_IN |= {"m/d/yy h:mm": 22, "[h]:mm:ss": 46, "h:mm:ss": 21}
Cell = tuple[str, str, str]
"""A stored cell: its kind, its value and its number format. The kinds are ``s`` (a shared
string), ``si`` (a shared string given as the XML of its item), ``is`` (an inline string),
``isx`` (an inline string given as XML), ``str`` (a text stored in the cell itself), ``n`` (a
number), ``d`` (a date written as ISO text), ``b`` (a truth value), ``e`` (an error value),
``f`` (a formula without a result), ``fn`` and ``fs`` (a formula with a number or a text as
its last result, written as ``formula|result``) and ``empty``."""


def typed(value: str) -> Cell:
    """Text typed into a text cell, as Excel and LibreOffice store it: a shared string."""
    return ("s", value, "@")


def number(value: str, number_format: str = "General") -> Cell:
    return ("n", value, number_format)


def serial(day: str) -> str:
    """The number a workbook stores for a date (days since 1899-12-30)."""
    return str((dt.date.fromisoformat(day) - dt.date(1899, 12, 30)).days)


def day_cell(day: str, number_format: str = "mm-dd-yy") -> Cell:
    return ("n", serial(day), number_format)


def text_xml(text: str) -> str:
    keep = ' xml:space="preserve"' if text != text.strip() or "\n" in text else ""
    return f"<t{keep}>{escape(text)}</t>"


def styles_xml() -> str:
    custom = [f for f in FORMATS if f not in BUILT_IN]
    ids = {f: 164 + n for n, f in enumerate(custom)} | BUILT_IN
    formats = "".join(
        f'<numFmt numFmtId="{ids[f]}" formatCode="{escape(f, {chr(34): "&quot;"})}"/>'
        for f in custom
    )
    xfs = "".join(
        f'<xf numFmtId="{ids[f]}" fontId="0" fillId="0" borderId="0" xfId="0"/>' for f in FORMATS
    )
    return (
        f'<styleSheet xmlns="{MAIN}"><numFmts count="{len(custom)}">{formats}</numFmts>'
        '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
        '<fills count="2"><fill><patternFill patternType="none"/></fill>'
        '<fill><patternFill patternType="gray125"/></fill></fills>'
        '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/>'
        f'</cellStyleXfs><cellXfs count="{len(FORMATS)}">{xfs}</cellXfs>'
        '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
        "</styleSheet>"
    )


class Stored:
    """A workbook written part by part, the way a spreadsheet program saves one."""

    def __init__(self, date1904: bool = False) -> None:
        self.strings: list[str] = []
        self.sheets: list[tuple[str, str, str]] = []
        self.date1904 = date1904

    def shared(self, item: str) -> int:
        if item not in self.strings:
            self.strings.append(item)
        return self.strings.index(item)

    def cell(self, ref: str, cell: Cell) -> str:
        kind, value, number_format = cell
        head = f'<c r="{ref}" s="{FORMATS.index(number_format)}"'
        if kind in ("s", "si"):
            index = self.shared(text_xml(value) if kind == "s" else value)
            return f'{head} t="s"><v>{index}</v></c>'
        if kind in ("is", "isx"):
            return f'{head} t="inlineStr"><is>{text_xml(value) if kind == "is" else value}</is></c>'
        if kind == "n":
            return f"{head}><v>{value}</v></c>"
        if kind in ("b", "e", "d", "str"):
            return f'{head} t="{kind}"><v>{escape(value)}</v></c>'
        if kind == "f":
            return f"{head}><f>{escape(value)}</f></c>"
        if kind in ("fn", "fs"):
            formula, _, result = value.partition("|")
            how = ' t="str"' if kind == "fs" else ""
            return f"{head}{how}><f>{escape(formula)}</f><v>{escape(result)}</v></c>"
        assert kind == "empty", kind
        return f"{head}/>"

    def add(
        self,
        title: str,
        cells: dict[tuple[int, int], Cell],
        *,
        hidden_rows: tuple[int, ...] = (),
        merged: tuple[str, ...] = (),
        auto_filter: str = "",
        state: str = "visible",
    ) -> None:
        rows = ""
        for r in sorted({r for r, _ in cells}):
            inner = "".join(
                self.cell(f"{X.column_letter(c)}{r}", cells[r, c])
                for c in sorted(c for row, c in cells if row == r)
            )
            hidden = ' hidden="1"' if r in hidden_rows else ""
            rows += f'<row r="{r}"{hidden}>{inner}</row>'
        tail = f'<autoFilter ref="{auto_filter}"/>' if auto_filter else ""
        if merged:
            ranges = "".join(f'<mergeCell ref="{span}"/>' for span in merged)
            tail += f'<mergeCells count="{len(merged)}">{ranges}</mergeCells>'
        xml = f'<worksheet xmlns="{MAIN}"><sheetData>{rows}</sheetData>{tail}</worksheet>'
        self.sheets.append((title, xml, state))

    def save(self, path: Path) -> Path:
        head = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        count = len(self.sheets)
        numbers = range(1, count + 1)
        types = (
            f'<Types xmlns="{PACKAGE}/content-types">'
            '<Default Extension="rels" '
            'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            f'<Override PartName="/xl/workbook.xml" ContentType="{SHEET_TYPE}.sheet.main+xml"/>'
            + "".join(
                f'<Override PartName="/xl/worksheets/sheet{n}.xml" '
                f'ContentType="{SHEET_TYPE}.worksheet+xml"/>'
                for n in numbers
            )
            + f'<Override PartName="/xl/styles.xml" ContentType="{SHEET_TYPE}.styles+xml"/>'
            f'<Override PartName="/xl/sharedStrings.xml" ContentType="{SHEET_TYPE}.sharedStrings+xml"/>'
            "</Types>"
        )
        package = (
            f'<Relationships xmlns="{PACKAGE}/relationships"><Relationship Id="rId1" '
            f'Type="{RELS}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
        )
        listed = "".join(
            f'<sheet name="{escape(title)}" sheetId="{n}" state="{state}" r:id="rId{n}"/>'
            for n, (title, _, state) in zip(numbers, self.sheets, strict=True)
        )
        epoch = ' date1904="1"' if self.date1904 else ""
        book = (
            f'<workbook xmlns="{MAIN}" xmlns:r="{RELS}"><workbookPr{epoch}/>'
            f"<sheets>{listed}</sheets></workbook>"
        )
        links = (
            f'<Relationships xmlns="{PACKAGE}/relationships">'
            + "".join(
                f'<Relationship Id="rId{n}" Type="{RELS}/worksheet" '
                f'Target="worksheets/sheet{n}.xml"/>'
                for n in numbers
            )
            + f'<Relationship Id="rId{count + 1}" Type="{RELS}/styles" Target="styles.xml"/>'
            f'<Relationship Id="rId{count + 2}" Type="{RELS}/sharedStrings" '
            'Target="sharedStrings.xml"/></Relationships>'
        )
        strings = (
            f'<sst xmlns="{MAIN}" count="{len(self.strings)}" uniqueCount="{len(self.strings)}">'
            + "".join(f"<si>{item}</si>" for item in self.strings)
            + "</sst>"
        )
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", head + types)
            archive.writestr("_rels/.rels", head + package)
            archive.writestr("xl/workbook.xml", head + book)
            archive.writestr("xl/_rels/workbook.xml.rels", head + links)
            archive.writestr("xl/styles.xml", head + styles_xml())
            archive.writestr("xl/sharedStrings.xml", head + strings)
            for n, (_, xml, _) in zip(numbers, self.sheets, strict=True):
                archive.writestr(f"xl/worksheets/sheet{n}.xml", head + xml)
        return path


def grid(blank: Path) -> dict[tuple[int, int], Cell]:
    """The cells of a blank sheet as a spreadsheet program saves the workbook made from it:
    a shared string in a text cell, an empty text cell where there is nothing."""
    base = X.read_csv_sheet(blank)
    cells: dict[tuple[int, int], Cell] = {}
    for r, row in enumerate(base.rows, start=1):
        for c in range(1, (len(base.header) if r > base.header_at else 1) + 1):
            text = row[c - 1] if c <= len(row) else ""
            cells[r, c] = typed(text) if text else ("empty", "", "@")
    return cells


def at(blank: Path, item: str, column: str) -> tuple[int, int]:
    """The worksheet row and column (from 1) of an item's cell, or of a ``#`` line's cell
    when ``item`` is the key of the line."""
    base = X.read_csv_sheet(blank)
    if column == "#":
        keys = [X.comment_key(row[0]) for row in base.rows[: base.header_at]]
        return keys.index(item) + 1, 1
    ids = [row[0] for row in base.rows]
    return ids.index(item, base.header_at) + 1, base.header.index(column) + 1


def saved(
    blank: Path,
    entries: dict[tuple[str, str], Cell] | None = None,
    name: str = "filled.xlsx",
    *,
    hidden_rows: tuple[int, ...] = (),
    merged: tuple[str, ...] = (),
    auto_filter: str = "",
) -> Path:
    """A workbook of the blank with entries (by item and column) as a program stores them."""
    cells = grid(blank)
    for (item, column), cell in (entries or {}).items():
        cells[at(blank, item, column)] = cell
    book = Stored()
    book.add("sheet1", cells, hidden_rows=hidden_rows, merged=merged, auto_filter=auto_filter)
    return book.save(blank.parent / name)


def one_cell(
    blank: Path, cell: Cell, column: str = "note"
) -> tuple[str | None, list[str], list[str]]:
    """What ``to-csv`` gives for the first item's cell of ``column`` stored as ``cell``: the
    text of the cell in the CSV (None when nothing is written), the problems and the notes."""
    base = X.read_csv_sheet(blank)
    first = base.rows[base.header_at + 1][0]
    text, problems, notes = X.to_csv(saved(blank, {(first, column): cell}), blank)
    return (parsed(text)[first][column] if text else None), problems, notes


# --------------------------------------------------------------------------------------------
# The round trip
# --------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "text", [literal_text(), outcome_text(), table_text()], ids=["literal", "outcome", "table"]
)
def test_a_workbook_from_a_blank_gives_back_its_bytes(tmp_path: Path, text: str) -> None:
    blank = written(tmp_path, "blank.csv", text)
    X.to_xlsx(blank, tmp_path / "blank.xlsx")
    back, problems, notes = X.to_csv(tmp_path / "blank.xlsx", blank)
    assert back.encode("utf-8") == blank.read_bytes()
    assert problems == [] and notes == []


@pytest.mark.parametrize("blank", REAL_BLANKS, ids=lambda path: path.stem)
def test_the_blank_sheets_of_the_study_round_trip(tmp_path: Path, blank: Path) -> None:
    if not blank.is_file():
        pytest.skip("the blank sheet is not on disk")
    assert X.main(["to-xlsx", str(blank), str(tmp_path / "sheet.xlsx")]) == 0
    code = X.main(
        [
            "to-csv",
            str(tmp_path / "sheet.xlsx"),
            "--blank",
            str(blank),
            "--out",
            str(tmp_path / "back.csv"),
        ]
    )
    assert code == 0
    assert (tmp_path / "back.csv").read_bytes() == blank.read_bytes()


def test_a_blank_with_windows_line_ends_round_trips(tmp_path: Path) -> None:
    items = tuple(item for item in LITERAL_ITEMS if "\n" not in "".join(item))
    blank = written(tmp_path, "blank.csv", literal_text(items, newline="\r\n"))
    assert X.read_csv_sheet(blank).newline == "\r\n"
    X.to_xlsx(blank, tmp_path / "blank.xlsx")
    back, problems, _ = X.to_csv(tmp_path / "blank.xlsx", blank)
    assert back.encode("utf-8") == blank.read_bytes() and not problems
    assert back.count("\r\n") == back.count("\n") == len(LITERAL_META) + 1 + len(items)


def test_the_cells_of_the_blank_are_the_cells_of_the_workbook(literal: Path) -> None:
    X.to_xlsx(literal, literal.with_suffix(".xlsx"))
    page = load_workbook(literal.with_suffix(".xlsx")).worksheets[0]
    base = X.read_csv_sheet(literal)
    for r, row in enumerate(base.rows, start=1):
        for c, text in enumerate(row, start=1):
            assert (page.cell(r, c).value or "") == text
    assert page.max_row == len(base.rows) and page.max_column == len(base.header)
    assert page["B19"].value == "'=cmd|' /C calc'!A0"  # the guard stays a character of the cell
    assert (page["C20"].value, page["C20"].data_type) == ("=SUM(1,2)", "s")  # text, no formula
    assert page["K18"].value.startswith("#1 reason") and page["J20"].value == LONG
    table = [cell for row in page.iter_rows(min_row=base.header_at + 1) for cell in row]
    assert {cell.number_format for cell in table if cell.coordinate != "J20"} == {"@"}
    assert {page.cell(r, 1).number_format for r in range(1, base.header_at + 1)} == {"@"}


def test_the_layout_of_the_blank_is_read_and_written(literal: Path, outcome: Path) -> None:
    quoted, bare = X.read_csv_sheet(literal), X.read_csv_sheet(outcome)
    assert (quoted.quoted_comments, quoted.header_at, quoted.newline) == (
        True,
        len(LITERAL_META),
        "\n",
    )
    assert (bare.quoted_comments, bare.header_at) == (False, len(OUTCOME_META))
    assert quoted.header == [*LITERAL_SHOWN, *LITERAL_ENTERED]
    assert bare.rows[6] == ["# session_start: "]  # the space after the colon is kept
    assert X.write_csv_sheet(quoted) == literal_text()
    lines = X.write_csv_sheet(bare).split("\n")
    assert lines[0] == "# task: B outcome audit (AUDIT_GUIDE.md section 6)"
    assert lines[len(OUTCOME_META)].startswith('"audit_id","generic",')
    assert all(
        line.startswith('"') and line.endswith('"') for line in lines[len(OUTCOME_META) : -1]
    )


def test_a_line_break_inside_a_cell_does_not_set_the_line_ends(tmp_path: Path) -> None:
    text = literal_text().replace("second line of the cell", "second\r\nline")
    assert X.read_csv_sheet(written(tmp_path, "blank.csv", text)).newline == "\n"


@pytest.mark.parametrize(
    ("text", "reason"),
    [
        (chr(0xFEFF) + literal_text(), "byte-order mark"),
        ('"# one","two"\n"a","b"\n', "a # line is not one cell"),
        ('"# only a line"\n', "no table"),
    ],
)
def test_a_file_that_is_not_a_blank_sheet_is_refused(
    tmp_path: Path, text: str, reason: str
) -> None:
    with pytest.raises(ValueError, match=reason):
        X.read_csv_sheet(written(tmp_path, "blank.csv", text))


def test_a_cell_too_long_for_a_worksheet_is_refused(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    items = (("I000000001", "x" * (X.TEXT_LIMIT + 1), *LITERAL_ITEMS[0][2:]),)
    blank = written(tmp_path, "blank.csv", literal_text(items))
    with pytest.raises(ValueError, match="longer than a worksheet cell"):
        X.to_xlsx(blank, tmp_path / "blank.xlsx")
    assert X.main(["to-xlsx", str(blank), str(tmp_path / "blank.xlsx")]) == 1
    assert "longer than a worksheet cell" in capsys.readouterr().err
    assert not (tmp_path / "blank.xlsx").exists()
    fits = (("I000000001", "x" * X.TEXT_LIMIT, *LITERAL_ITEMS[0][2:]),)
    blank = written(tmp_path, "fits.csv", literal_text(fits))
    assert X.main(["to-xlsx", str(blank), str(tmp_path / "fits.xlsx")]) == 0


def test_a_workbook_that_does_not_give_back_the_blank_is_not_kept(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    # a carriage return inside a cell does not survive a workbook: the sheet is refused
    text = literal_text().replace("second line of the cell", "second\rline")
    blank = written(tmp_path, "blank.csv", text)
    assert X.main(["to-xlsx", str(blank), str(tmp_path / "blank.xlsx")]) == 1
    assert "does not give back the blank sheet" in capsys.readouterr().err
    assert not (tmp_path / "blank.xlsx").exists()


# --------------------------------------------------------------------------------------------
# The workbook that to-xlsx writes
# --------------------------------------------------------------------------------------------


def sheet_xml(path: Path) -> ElementTree.Element:
    with zipfile.ZipFile(path) as archive:
        return ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))


def text_cells(path: Path) -> tuple[dict[str, str], dict[str, str]]:
    """The number format id and the kind of every cell of the first worksheet, by reference."""
    with zipfile.ZipFile(path) as archive:
        styles = ElementTree.fromstring(archive.read("xl/styles.xml"))
    cell_formats = styles.find(f"{{{MAIN}}}cellXfs")
    assert cell_formats is not None
    formats = [xf.get("numFmtId", "") for xf in cell_formats]
    found = {c.get("r", ""): c for c in sheet_xml(path).iter(f"{{{MAIN}}}c")}
    return (
        {ref: formats[int(c.get("s", "0"))] for ref, c in found.items()},
        {ref: c.get("t", "") for ref, c in found.items() if len(c)},
    )


def cells_to_fill(blank: Path) -> list[str]:
    """The cells an annotator fills: the entered columns of the items, and the ``#`` lines
    that name no value."""
    base = X.read_csv_sheet(blank)
    columns = range(X.first_entered(base.header) + 1, len(base.header) + 1)
    rows = range(base.header_at + 2, len(base.rows) + 1)
    lines = [
        n for n, row in enumerate(base.rows[: base.header_at], 1) if not X.comment_value(row[0])
    ]
    return [f"{X.column_letter(c)}{r}" for r in rows for c in columns] + [f"A{n}" for n in lines]


def lists_of(path: Path) -> dict[str, dict[str, str]]:
    """The list of each column that has one: its range, its values and its attributes."""
    rules = sheet_xml(path).find(f"{{{MAIN}}}dataValidations")
    found = {}
    for rule in rules if rules is not None else ():
        formula = rule.find(f"{{{MAIN}}}formula1")
        assert formula is not None and formula.text
        column = re.match(r"[A-Z]+", rule.get("sqref", ""))
        assert column is not None
        found[column.group(0)] = {**rule.attrib, "values": formula.text}
    return found


@pytest.mark.parametrize("which", ["literal", "outcome"])
def test_every_cell_of_the_workbook_is_text(which: str, request: pytest.FixtureRequest) -> None:
    blank = request.getfixturevalue(which)
    X.to_xlsx(blank, blank.with_suffix(".xlsx"))
    formats, kinds = text_cells(blank.with_suffix(".xlsx"))
    # the format "@": what is typed stays what was typed; a shown cell too long for every
    # version of Excel to show in that format keeps the general one
    long_cells = {"J20"} if which == "literal" else set()
    assert {ref for ref, number_format in formats.items() if number_format != "49"} == long_cells
    assert {formats[ref] for ref in long_cells} <= {"0"}
    assert set(kinds.values()) == {"inlineStr"}
    to_fill = cells_to_fill(blank)
    assert len(to_fill) == (4 * 11 + 2 if which == "literal" else 3 * 13 + 2)
    assert all(formats[ref] == "49" for ref in to_fill)
    assert not [ref for ref in to_fill if ref in kinds and not ref.startswith("A")]


def test_a_long_shown_cell_is_not_in_the_text_format(tmp_path: Path) -> None:
    long, longer = "x" * X.TEXT_SHOWN, "x" * (X.TEXT_SHOWN + 1)
    items = (("I000000001", long, longer, *LITERAL_ITEMS[0][3:]),)
    blank = written(tmp_path, "blank.csv", literal_text(items))
    assert X.main(["to-xlsx", str(blank), str(tmp_path / "blank.xlsx")]) == 0
    formats, kinds = text_cells(tmp_path / "blank.xlsx")
    assert (formats["B17"], formats["C17"]) == ("49", "0")
    assert kinds["B17"] == kinds["C17"] == "inlineStr"  # text in both


def test_the_lists_offer_the_closed_sets_and_refuse_nothing(literal: Path, outcome: Path) -> None:
    X.to_xlsx(literal, literal.with_suffix(".xlsx"))
    X.to_xlsx(outcome, outcome.with_suffix(".xlsx"))
    lists = lists_of(literal.with_suffix(".xlsx"))
    assert {k: v["values"] for k, v in lists.items()} == {
        "L": '"recovery,next_delivery,depletion,discontinuation,none"',
        "O": '"1,0"',
        "P": '"tbd,no_date,vague,no_statement"',
        "Q": '"asserted,estimated,undetermined,no_statement"',
        "U": '"1,0"',
    }
    first, last = len(LITERAL_META) + 2, len(LITERAL_META) + 1 + len(LITERAL_ITEMS)
    assert {k: v["sqref"] for k, v in lists.items()} == {
        k: f"{k}{first}:{k}{last}" for k in "LOPQU"
    }
    lists |= lists_of(outcome.with_suffix(".xlsx"))
    assert {k: lists[k]["values"] for k in ("AC", "AD", "AG", "AJ")} == {
        "AC": '"ok,error,cannot_tell"',
        "AD": '"ok,error,cannot_tell"',
        "AG": '"recovered,discontinued,censored,not_at_risk"',
        "AJ": '"recovered,discontinued,censored,not_at_risk"',
    }
    for rule in lists.values():
        assert rule["type"] == "list" and rule["allowBlank"] == "1"
        assert rule["showErrorMessage"] == "0"  # no message, so no entry is ever refused
        assert rule["showDropDown"] == "0"  # which, in this format, shows the arrow


def study_constants(module: str, names: set[str]) -> dict[str, tuple[str, ...]]:
    """Tuples of a module of the study, read from its source (the module itself needs the
    study's environment)."""
    tree = ast.parse((HERE / module).read_text(encoding="utf-8"))
    found = {}
    for node in tree.body:
        target = node.targets[0] if isinstance(node, ast.Assign) else None
        if isinstance(target, ast.Name) and target.id in names:
            found[target.id] = ast.literal_eval(node.value)
    assert set(found) == names
    return found


def test_the_lists_hold_what_the_validators_accept(literal: Path, outcome: Path) -> None:
    sample = study_constants(
        "audit_sample.py", {"STATEMENT_TYPES", "CERTAINTY", "ABSTAIN_REASONS", "SHOWN", "ENTERED"}
    )
    outcomes = study_constants("audit_outcomes.py", {"TYPES", "VERDICTS", "SHOWN", "ENTERED"})
    assert (*sample["SHOWN"], *sample["ENTERED"]) == (*LITERAL_SHOWN, *LITERAL_ENTERED)
    assert (*outcomes["SHOWN"], *outcomes["ENTERED"]) == (*OUTCOME_SHOWN, *OUTCOME_ENTERED)
    lists = X.value_lists(X.read_csv_sheet(literal))
    assert {name: set(values) for name, values in lists.items()} == {
        "statement_type": set(sample["STATEMENT_TYPES"]),
        "abstain": {"0", "1"},
        "abstain_reason": set(sample["ABSTAIN_REASONS"]),
        "certainty": set(sample["CERTAINTY"]),
        "hard": {"0", "1"},
    }
    lists = X.value_lists(X.read_csv_sheet(outcome))
    assert {name: set(values) for name, values in lists.items()} == {
        "verdict_B": set(outcomes["VERDICTS"]),
        "verdict_A": set(outcomes["VERDICTS"]),
        "true_B_type": set(outcomes["TYPES"]),
        "true_A_type": set(outcomes["TYPES"]),
    }
    assert X.first_entered([*sample["SHOWN"], *sample["ENTERED"]]) == len(sample["SHOWN"])
    assert X.first_entered([*outcomes["SHOWN"], *outcomes["ENTERED"]]) == len(outcomes["SHOWN"])
    assert X.first_entered(["audit_id", "capture", "status"]) == 3


def test_frozen_rows_and_folded_columns(literal: Path, outcome: Path, tmp_path: Path) -> None:
    X.to_xlsx(literal, literal.with_suffix(".xlsx"))
    X.to_xlsx(outcome, outcome.with_suffix(".xlsx"))
    page = load_workbook(literal.with_suffix(".xlsx")).worksheets[0]
    assert page.freeze_panes == f"B{len(LITERAL_META) + 2}"
    assert not any(d.outlineLevel for d in page.column_dimensions.values())  # 10 shown columns
    book = load_workbook(outcome.with_suffix(".xlsx"))
    page = book.worksheets[0]
    assert page.freeze_panes == f"B{len(OUTCOME_META) + 2}"
    grouped = [d for d in page.column_dimensions.values() if d.outlineLevel]
    columns = sorted(c for d in grouped for c in range(d.min, d.max + 1))
    assert columns == list(range(2, len(OUTCOME_SHOWN) + 1))  # B to AB: not the id
    # the annotator folds the group, unfreezes the rows and saves: the same sheet is read
    for dimension in grouped:
        dimension.hidden = True
    page.freeze_panes = None
    book.save(tmp_path / "folded.xlsx")
    back, problems, notes = X.to_csv(tmp_path / "folded.xlsx", outcome)
    assert back.encode("utf-8") == outcome.read_bytes() and not problems and not notes


def test_the_workbook_names_nobody(literal: Path) -> None:
    X.to_xlsx(literal, literal.with_suffix(".xlsx"))
    with zipfile.ZipFile(literal.with_suffix(".xlsx")) as archive:
        core = archive.read("docProps/core.xml").decode("utf-8")
        parts = set(archive.namelist())
    assert parts == {
        "[Content_Types].xml",
        "_rels/.rels",
        "docProps/app.xml",
        "docProps/core.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/styles.xml",
        "xl/theme/theme1.xml",
        "xl/workbook.xml",
        "xl/worksheets/sheet1.xml",
    }
    people = dict(re.findall(r"<(?:dc|cp):(creator|lastModifiedBy)[^>]*>([^<]*)</", core))
    # an element written empty or left out names nobody (the library does either)
    assert set(people) <= {"creator", "lastModifiedBy"} and set(people.values()) <= {""}


# --------------------------------------------------------------------------------------------
# Cells as a spreadsheet program stores them
# --------------------------------------------------------------------------------------------


@pytest.mark.parametrize("which", ["literal", "outcome"])
def test_a_workbook_saved_by_a_program_gives_back_the_blank(
    which: str, request: pytest.FixtureRequest
) -> None:
    blank = request.getfixturevalue(which)
    back, problems, notes = X.to_csv(saved(blank), blank)  # shared strings, empty text cells
    assert back.encode("utf-8") == blank.read_bytes()
    assert problems == [] and notes == []


TYPED = (
    "a, b",
    'say "this"',
    "one; two | three",
    "first line\nsecond line",
    "'apostrophe first",
    "=1+1",
    "+1 lot",
    "-none-",
    "@home",
    "  spaces around  ",
    "\u00a0no-break spaces\u00a0",
    "en dash – and ’ and 供給",
    "#not a line",
    "TRUE",
    "2021-03",
    "0012",
    "1e5",
    "09:30",
    LONG,
)
"""Entries that CSV has to quote, or that a spreadsheet program would convert in a cell that
is not text."""


@pytest.mark.parametrize("kind", ["s", "is", "str"])
@pytest.mark.parametrize("entry", TYPED, ids=range(len(TYPED)))
def test_text_typed_into_a_text_cell_is_written_as_typed(
    literal: Path, entry: str, kind: str
) -> None:
    text, problems, notes = one_cell(literal, (kind, entry, "@"))
    assert text == entry
    assert problems == [] and notes == []


def test_every_entered_cell_reaches_the_csv(literal: Path) -> None:
    entries = {
        (item[0], column): typed(f"{item[0]} {column} {TYPED[(n + m) % 13]}")
        for n, item in enumerate(LITERAL_ITEMS)
        for m, column in enumerate(LITERAL_ENTERED)
    }
    text, problems, notes = X.to_csv(saved(literal, entries), literal)
    assert problems == [] and notes == []
    table = parsed(text)
    assert {(i, c): typed(table[i][c]) for i in table for c in LITERAL_ENTERED} == entries
    shown = parsed(literal_text())
    assert all(table[i][c] == shown[i][c] for i in table for c in LITERAL_SHOWN)
    assert (
        text.split("\n")[: len(LITERAL_META) + 1]
        == literal_text().split("\n")[: len(LITERAL_META) + 1]
    )
    # every cell of the table is quoted, whatever it holds
    assert text.startswith('"# task')
    rows = text.split("\n")[len(LITERAL_META)].split('","')
    assert len(rows) == len(LITERAL_SHOWN) + len(LITERAL_ENTERED)


def test_runs_of_formatted_text_are_one_text(literal: Path) -> None:
    runs = (
        '<r><rPr><b/><sz val="11"/></rPr><t>by </t></r><r><t xml:space="preserve">mid June </t></r>'
        "<r><rPr><i/></rPr><t>2021</t></r>"
        '<rPh sb="0" eb="2"><t>カナ</t></rPh><phoneticPr fontId="1"/>'
    )
    assert one_cell(literal, ("si", runs, "@"), "quote") == ("by mid June 2021", [], [])
    assert one_cell(literal, ("isx", runs, "@"), "quote") == ("by mid June 2021", [], [])


@pytest.mark.parametrize(
    ("cell", "text", "stored"),
    [
        (number("1"), "1", "a number (shown as General)"),
        (number("0"), "0", "a number (shown as General)"),
        (number("1.0"), "1", "a number (shown as General)"),
        (number("0.5", "0%"), "0.5", "a number (shown as 0%)"),
        (number("-3"), "-3", "a number (shown as General)"),
        (number("7", "@"), "7", "a number (shown as @)"),
        (number("20261001"), "20261001", "a number (shown as General)"),
        (("b", "1", "General"), "TRUE", "a truth value"),
        (("b", "0", "General"), "FALSE", "a truth value"),
        (day_cell("2021-06-15"), "2021-06-15", "a date (shown as mm-dd-yy)"),
        (day_cell("2021-06-01"), "2021-06-01", "a date (shown as mm-dd-yy)"),
        (
            day_cell("2021-06-15", "yyyy\\-mm\\-dd"),
            "2021-06-15",
            "a date (shown as yyyy\\-mm\\-dd)",
        ),
        (day_cell("2021-06-15", "d-mmm"), "2021-06-15", "a date (shown as d-mmm)"),
        (day_cell("2021-06-01", "mmm-yy"), "2021-06", "a date (shown as mmm-yy)"),
        (
            day_cell("2021-12-01", "[$-409]mmmm\\ yyyy;@"),
            "2021-12",
            "a date (shown as [$-409]mmmm\\ yyyy;@)",
        ),
        (day_cell("2021-06-15", "mmm-yy"), "2021-06-15", "a date (shown as mmm-yy)"),
        (day_cell("2021-06-01", FORMATS[8]), "2021-06", f"a date (shown as {FORMATS[8]})"),
        (day_cell("2021-06-01", FORMATS[11]), "2021-06-01", f"a date (shown as {FORMATS[11]})"),
        (number("44362.5", "m/d/yy h:mm"), "2021-06-15 12:00", "a date (shown as m/d/yy h:mm)"),
        (number("0.3958333333", "h:mm"), "09:30", "a time (shown as h:mm)"),
        (number("0.3960069444", "h:mm:ss"), "09:30:15", "a time (shown as h:mm:ss)"),
        (number("1.5", "[h]:mm:ss"), "1 day, 12:00:00", "a duration (shown as [h]:mm:ss)"),
        (("d", "2021-06-01T00:00:00", "General"), "2021-06-01", "a date (shown as General)"),
        (("d", "2021-06-15", "mm-dd-yy"), "2021-06-15", "a date (shown as mm-dd-yy)"),
    ],
)
def test_a_cell_that_is_not_text_is_written_as_text_and_named(
    literal: Path, cell: Cell, text: str, stored: str
) -> None:
    got, problems, notes = one_cell(literal, cell, "end")
    assert got == text
    assert problems == []
    assert notes == [f"N17: stored as {stored}, written as the text {text!r}"]


def test_a_month_is_not_written_as_its_first_day() -> None:
    first = dt.datetime(2021, 6, 1)
    assert X.stored_text(first, "mmm-yy")[0] == "2021-06"
    assert X.stored_text(first, "yyyy-mm")[0] == "2021-06"
    assert X.stored_text(first, 'yyyy"年"m"月";@')[0] == "2021-06"
    assert X.stored_text(first, 'mmm\\ yyyy "d"')[0] == "2021-06"  # a quoted d is no day
    assert X.stored_text(first, "[$-F800]mmmm yyyy")[0] == "2021-06"  # nor one in brackets
    assert X.stored_text(first, "mmm\\-yy;d")[0] == "2021-06"  # nor one of a later section
    assert X.stored_text(first, "yyyy-mm-dd")[0] == "2021-06-01"
    assert X.stored_text(first, "d-mmm-yy")[0] == "2021-06-01"
    assert X.stored_text(first, "yyyy")[0] == "2021-06-01"  # no month shown: the day stored
    assert X.stored_text(first, "h:mm")[0] == "2021-06-01"  # minutes, not a month
    assert X.stored_text(first, "mm:ss")[0] == "2021-06-01"
    assert X.stored_text(first, "General")[0] == "2021-06-01"
    assert X.stored_text(dt.datetime(2021, 6, 2), "mmm-yy")[0] == "2021-06-02"
    assert X.stored_text(dt.date(2021, 6, 1), "mmm-yy")[0] == "2021-06"
    assert X.stored_text(dt.datetime(2021, 6, 1, 9, 30), "mmm-yy")[0] == "2021-06-01 09:30"
    assert X.stored_text(dt.time(9, 30), "h:mm") == ("09:30", "a time (shown as h:mm)")
    assert X.stored_text(True, "General") == ("TRUE", "a truth value")
    assert X.stored_text(1, "General")[0] == X.stored_text(1.0, "General")[0] == "1"
    assert X.stored_text(2.5, "General")[0] == "2.5"


def test_a_workbook_that_counts_days_from_1904(literal: Path) -> None:
    cells = grid(literal)
    days = (dt.date(2021, 6, 15) - dt.date(1904, 1, 1)).days
    cells[at(literal, "I000000001", "start")] = number(str(days), "mm-dd-yy")
    book = Stored(date1904=True)
    book.add("sheet1", cells)
    text, problems, _ = X.to_csv(book.save(literal.parent / "mac.xlsx"), literal)
    assert parsed(text)["I000000001"]["start"] == "2021-06-15" and not problems


@pytest.mark.parametrize(
    ("cell", "shown"),
    [
        (("f", "recovery", "General"), "=recovery"),  # typed =recovery; never calculated
        (("fn", "1+0|1", "General"), "=1+0"),
        (("fs", '"reco"&"very"|recovery', "General"), '="reco"&"very"'),
        (("f", "+1", "@"), "=+1"),
    ],
)
def test_a_formula_stops_the_conversion(literal: Path, cell: Cell, shown: str) -> None:
    text, problems, notes = one_cell(literal, cell, "statement_type")
    assert text is None  # nothing is written: not the result, and not an empty label
    assert len(problems) == 1 and problems[0].startswith(
        "L17: holds a formula, which is not an entry"
    )
    assert problems[0].endswith(f": {shown!r}")
    assert notes == []


@pytest.mark.parametrize("value", ["#NAME?", "#DIV/0!", "#N/A", "#VALUE!"])
def test_an_error_value_stops_the_conversion(literal: Path, value: str) -> None:
    text, problems, _ = one_cell(literal, ("e", value, "General"))
    assert text is None
    assert problems == [f"V17: holds an error value, which is not an entry: {value}"]
    text, problems, _ = one_cell(literal, ("fn", f"nothing|{value}", "General"))
    assert text is None and "holds a formula" in problems[0]


def test_a_carriage_return_inside_a_cell_becomes_a_line_break(literal: Path) -> None:
    # pasted text with Windows line ends: Excel writes the carriage return as _x000D_
    text, problems, notes = one_cell(literal, typed("line one_x000D_\nline two"))
    assert text == "line one\nline two"
    assert problems == []
    assert notes == [
        "V17: escaped characters are written as the characters they name",
        "V17: a carriage return inside the cell is written as a line break",
    ]
    # LibreOffice writes it as a character reference
    text, problems, notes = one_cell(literal, ("isx", "<t>one&#13;\ntwo&#13;three</t>", "@"))
    assert text == "one\ntwo\nthree"
    assert notes == ["V17: a carriage return inside the cell is written as a line break"]
    # the CSV keeps one kind of line end, also for whoever reads it as a blank
    back, _, _ = X.to_csv(saved(literal, {("I000000001", "note"): typed("a_x000D_\nb")}), literal)
    assert "\r" not in back
    assert X.read_csv_sheet(written(literal.parent, "back.csv", back)).newline == "\n"


def test_other_escaped_characters_are_read(literal: Path) -> None:
    text, _, notes = one_cell(literal, typed("tab_x0009_here, vertical_x000B_tab"))
    assert text == "tab\there, vertical\x0btab"
    assert notes == ["V17: escaped characters are written as the characters they name"]
    # an escaped underscore before text of that form: what was typed cannot be told any more
    text, problems, _ = one_cell(literal, typed("typed _x005F_x000D_ as text"))
    assert text is None
    assert len(problems) == 1 and "cannot be read with certainty" in problems[0]
    # that escape alone is read, and puts no other cell in doubt
    entries = {
        ("I000000001", "note"): typed("an _x005F_ alone"),
        ("I000000002", "note"): typed("an_x_ and _x12_ and _xD_"),
    }
    text, problems, notes = X.to_csv(saved(literal, entries), literal)
    assert problems == [] and notes == []
    assert parsed(text)["I000000001"]["note"] == "an _ alone"
    assert parsed(text)["I000000002"]["note"] == "an_x_ and _x12_ and _xD_"


def test_spaces_typed_by_the_annotator_are_kept(literal: Path) -> None:
    entries = {
        ("I000000001", "statement_type"): typed("recovery "),
        ("I000000001", "abstain_reason"): typed("\u00a0vague"),
        ("I000000001", "note"): typed(" "),
        ("I000000001", "quote"): typed("\u3000mid\u00a0June\t"),
    }
    text, problems, notes = X.to_csv(saved(literal, entries), literal)
    row = parsed(text)["I000000001"]
    assert (row["statement_type"], row["abstain_reason"]) == ("recovery ", "\u00a0vague")
    assert (row["note"], row["quote"]) == (" ", "\u3000mid\u00a0June\t")
    assert problems == [] and notes == []  # the validators strip the cells; nothing is changed


def test_a_comment_on_a_cell_is_named(literal: Path) -> None:
    X.to_xlsx(literal, literal.with_suffix(".xlsx"))
    book = load_workbook(literal.with_suffix(".xlsx"))
    book.worksheets[0]["L17"].comment = Comment("not sure about this one", "")
    book.save(literal.parent / "commented.xlsx")
    text, problems, notes = X.to_csv(literal.parent / "commented.xlsx", literal)
    assert text.encode("utf-8") == literal.read_bytes() and problems == []
    assert notes == ["L17: carries a comment, which is not written: 'not sure about this one'"]


# --------------------------------------------------------------------------------------------
# Worksheets, rows and columns
# --------------------------------------------------------------------------------------------

ANSWER = {("I000000001", "statement_type"): typed("recovery")}


def with_sheets(blank: Path, order: tuple[str, ...], **titles: dict[tuple[int, int], Cell]) -> Path:
    book = Stored()
    for title in order:
        book.add(title, titles[title])
    return book.save(blank.parent / "sheets.xlsx")


def test_a_renamed_or_a_hidden_worksheet_is_read(literal: Path) -> None:
    cells = grid(literal) | {at(literal, "I000000001", "statement_type"): typed("recovery")}
    book = Stored()
    book.add("pilot (my copy) – final", cells, state="hidden")
    text, problems, notes = X.to_csv(book.save(literal.parent / "renamed.xlsx"), literal)
    assert parsed(text)["I000000001"]["statement_type"] == "recovery"
    assert problems == [] and notes == []


def test_a_second_worksheet_is_named_and_not_read(literal: Path) -> None:
    cells = grid(literal) | {at(literal, "I000000001", "statement_type"): typed("recovery")}
    scrap = {(1, 1): typed("my notes"), (2, 2): typed("ask about row 3"), (3, 1): typed(" ")}
    path = with_sheets(
        literal, ("sheet1", "Sheet2", "Sheet3"), sheet1=cells, Sheet2=scrap, Sheet3={}
    )
    text, problems, notes = X.to_csv(path, literal)
    assert parsed(text)["I000000001"]["statement_type"] == "recovery" and problems == []
    assert notes == ["worksheet 'Sheet2' is not read (2 cells with text)"]


def test_the_table_is_found_when_another_worksheet_comes_first(literal: Path) -> None:
    # Numbers can put a summary worksheet before the tables it exports
    cells = grid(literal) | {at(literal, "I000000001", "statement_type"): typed("recovery")}
    summary = {(1, 1): typed("This document was exported from Numbers.")}
    path = with_sheets(
        literal, ("Export Summary", "sheet1"), **{"Export Summary": summary, "sheet1": cells}
    )
    text, problems, notes = X.to_csv(path, literal)
    assert parsed(text)["I000000001"]["statement_type"] == "recovery" and problems == []
    assert notes == [
        "the sheet is read from worksheet 'sheet1', which is not the first",
        "worksheet 'Export Summary' is not read (1 cells with text)",
    ]


def test_two_worksheets_with_the_table_are_not_chosen_between(
    literal: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    # a copy of the worksheet, made as a backup: which one holds the labels is not guessed
    old = grid(literal) | {at(literal, "I000000001", "statement_type"): typed("none")}
    new = grid(literal) | {at(literal, "I000000001", "statement_type"): typed("recovery")}
    path = with_sheets(literal, ("sheet1 (2)", "sheet1"), **{"sheet1 (2)": old, "sheet1": new})
    text, problems, _ = X.to_csv(path, literal)
    assert text == ""
    assert problems == [
        "more than one worksheet holds the table ('sheet1 (2)', 'sheet1'); name one with --sheet"
    ]
    for title, label in (("sheet1", "recovery"), ("sheet1 (2)", "none")):
        text, problems, notes = X.to_csv(path, literal, sheet=title)
        assert parsed(text)["I000000001"]["statement_type"] == label and problems == []
        assert [n for n in notes if "is not read" in n] != []
    assert X.to_csv(path, literal, sheet="Sheet9") == (
        "",
        ["no worksheet named 'Sheet9'; the workbook holds 'sheet1 (2)', 'sheet1'"],
        [],
    )
    out = literal.parent / "chosen.csv"
    arguments = ["to-csv", str(path), "--blank", str(literal), "--out", str(out)]
    assert X.main(arguments) == 1 and not out.exists()
    assert X.main([*arguments, "--sheet", "sheet1"]) == 0
    assert parsed(out.read_text(encoding="utf-8"))["I000000001"]["statement_type"] == "recovery"
    capsys.readouterr()


def test_hidden_and_filtered_rows_are_read(literal: Path) -> None:
    entries = {(item[0], "statement_type"): typed("none") for item in LITERAL_ITEMS}
    first = len(LITERAL_META) + 2
    path = saved(
        literal, entries, hidden_rows=(first, first + 2), auto_filter=f"A{first - 1}:V{first + 3}"
    )
    text, problems, notes = X.to_csv(path, literal)
    assert [row["statement_type"] for row in parsed(text).values()] == ["none"] * 4
    assert problems == [] and notes == []


Rows = list[list[Cell | None]]
Change = Callable[[Rows, int], Rows]
"""What an annotator does to the rows of the worksheet; the number is the index of the header
row."""


def rows_of(blank: Path, entries: dict[tuple[str, str], Cell]) -> Rows:
    """The stored cells of a filled workbook, row by row."""
    cells = grid(blank) | {
        at(blank, item, column): cell for (item, column), cell in entries.items()
    }
    height, width = max(r for r, _ in cells), max(c for _, c in cells)
    return [[cells.get((r, c)) for c in range(1, width + 1)] for r in range(1, height + 1)]


def from_rows(blank: Path, rows: Rows, name: str = "moved.xlsx") -> Path:
    book = Stored()
    book.add(
        "sheet1",
        {(r, c): cell for r, row in enumerate(rows, 1) for c, cell in enumerate(row, 1) if cell},
    )
    return book.save(blank.parent / name)


def moved(blank: Path, change: Change) -> tuple[str, list[str], list[str]]:
    """``to-csv`` on a workbook with one answer whose rows were changed."""
    rows = change(rows_of(blank, ANSWER), X.read_csv_sheet(blank).header_at)
    return X.to_csv(from_rows(blank, rows), blank)


def test_sorted_rows_stop_the_conversion(literal: Path) -> None:
    text, problems, _ = moved(literal, lambda rows, h: rows[: h + 1] + rows[h + 1 :][::-1])
    assert text == ""
    assert problems == [
        "the item ids, or their order, are not the blank's: the rows are in another order "
        "(row 17 holds 'I000000004', the blank 'I000000001')"
    ]


def test_empty_rows_between_the_items_are_left_out(literal: Path) -> None:
    spaces = [None, None, None, None, None, typed("  ")]
    text, problems, notes = moved(
        literal, lambda rows, h: [*rows[:h], [], *rows[h : h + 2], [], spaces, *rows[h + 2 :]]
    )
    assert problems == []
    assert notes == ["2 empty rows between the rows of the table are left out"]
    expected = literal_text().replace('"I000000001"', "@", 1).split("@")
    assert (
        text.startswith(expected[0]) and parsed(text)["I000000001"]["statement_type"] == "recovery"
    )
    assert len(text.split("\n")) == len(literal_text().split("\n"))


@pytest.mark.parametrize(
    ("change", "said"),
    [
        (lambda rows, h: rows[: h + 2] + rows[h + 3 :], "'I000000002' is missing"),
        (lambda rows, h: rows[: h + 3] + rows[h + 2 :], "'I000000002' is there 2 times"),
        (lambda rows, h: [*rows, [typed("I000000009")]], "row 21 holds 'I000000009', not an item"),
        (lambda rows, h: [*rows, [None, None, typed("see above")]], "row 21 holds '', not an item"),
        (
            lambda rows, h: [*rows, [], [], [typed("# done at 10:00")]],
            "row 23 holds '# done at 10:00'",
        ),
        (
            lambda rows, h: [
                *rows[: h + 1],
                [typed(" I000000001"), *rows[h + 1][1:]],
                *rows[h + 2 :],
            ],
            "'I000000001' is missing; row 17 holds ' I000000001', not an item",
        ),
        (
            lambda rows, h: [*rows[: h + 2], rows[h], *rows[h + 2 :]],
            "row 18 holds 'item_id', not an item",
        ),
    ],
    ids=["deleted", "twice", "added", "text below", "line below", "id changed", "header twice"],
)
def test_changed_items_stop_the_conversion(literal: Path, change: Change, said: str) -> None:
    text, problems, _ = moved(literal, change)
    assert text == ""
    assert len(problems) == 1 and problems[0].startswith(
        "the item ids, or their order, are not the blank's: "
    )
    assert said in problems[0]


@pytest.mark.parametrize(
    ("change", "said"),
    [
        (
            lambda rows, h: [[*r[:12], None, *r[12:]] if n >= h else r for n, r in enumerate(rows)],
            "M holds '' for 'start'; N holds 'start' for 'end'; O holds 'end' for 'abstain'; and 8 more",
        ),
        (
            lambda rows, h: [r[:12] + r[13:] if n >= h else r for n, r in enumerate(rows)],
            "M holds 'end' for 'start'",
        ),
        (lambda rows, h: [[None, *r] for r in rows], "A holds '' for 'item_id'"),
        (
            lambda rows, h: [
                *rows[:h],
                [*rows[h][:11], typed("type"), *rows[h][12:]],
                *rows[h + 1 :],
            ],
            "L holds 'type' for 'statement_type'",
        ),
        (lambda rows, h: rows[:h] + rows[h + 1 :], "A holds 'I000000001' for 'item_id'"),
        (
            lambda rows, h: [*rows[:h], [*rows[h][:21], typed("notes")], *rows[h + 1 :]],
            "V holds 'notes' for 'note'",
        ),
        (
            lambda rows, h: [*rows[:h], [typed("Item_id"), *rows[h][1:]], *rows[h + 1 :]],
            "A holds 'Item_id' for 'item_id'",
        ),
    ],
    ids=[
        "column inserted",
        "column deleted",
        "column before the first",
        "name changed",
        "row deleted",
        "last name changed",
        "first name changed",
    ],
)
def test_a_changed_header_stops_the_conversion(literal: Path, change: Change, said: str) -> None:
    text, problems, _ = moved(literal, change)
    assert text == ""
    assert len(problems) == 1 and problems[0].startswith(
        "the header row of the table is not the blank's: "
    )
    assert said in problems[0]


def test_a_workbook_without_a_table_is_refused(literal: Path) -> None:
    text, problems, _ = moved(literal, lambda rows, h: rows[:h])
    assert (text, problems) == ("", ["no table under the # lines"])


def test_text_to_the_right_of_the_last_column_stops_the_conversion(literal: Path) -> None:
    beside = [typed("my remark: unsure about this one"), None, typed(" "), typed("another")]
    text, problems, notes = moved(
        literal, lambda rows, h: [*rows[: h + 1], rows[h + 1] + beside, *rows[h + 2 :]]
    )
    assert text == ""
    assert problems == [
        "W17: text to the right of the last column: 'my remark: unsure about this one'",
        "Z17: text to the right of the last column: 'another'",
    ]
    assert notes == ["Y17: a cell of spaces to the right of the last column is left out"]
    _, problems, _ = moved(
        literal, lambda rows, h: [*rows[:h], [*rows[h], typed("my column")], *rows[h + 1 :]]
    )
    assert problems == ["W16: text to the right of the last column: 'my column'"]


def test_spaces_to_the_right_of_the_last_column_do_not_stop_it(literal: Path) -> None:
    text, problems, notes = moved(
        literal, lambda rows, h: [*rows[: h + 1], [*rows[h + 1], typed("  ")], *rows[h + 2 :]]
    )
    assert problems == [] and parsed(text)["I000000001"]["statement_type"] == "recovery"
    assert notes == ["W17: a cell of spaces to the right of the last column is left out"]


def test_merged_cells(literal: Path) -> None:
    entries = {("I000000001", "certainty"): typed("asserted")}
    text, problems, notes = X.to_csv(saved(literal, entries, merged=("Q17:Q18",)), literal)
    assert problems == [
        "merged cells Q17:Q18 in the columns to fill: the cells under a merge are empty, so "
        "each of them has to be filled on its own"
    ]
    # across columns, and reaching into the columns to fill from the shown ones
    for span in ("L17:M17", "K17:L17", "V20:X20", "L16:L17"):
        assert len(X.to_csv(saved(literal, entries, merged=(span,)), literal)[1]) == 1, span
    # a merge among the # lines or in the shown columns loses no entry
    text, problems, notes = X.to_csv(saved(literal, entries, merged=("A1:F1", "B17:B18")), literal)
    assert problems == []
    assert notes == [
        "shown cell changed: item I000000002, column drug (B18): '', "
        "the blank holds 'Sulfasalazine Tablets'",
        "merged cells A1:F1: only the first cell of a merge holds text",
        "merged cells B17:B18: only the first cell of a merge holds text",
    ]
    assert parsed(text)["I000000001"]["certainty"] == "asserted"
    # nor does one that widens the last column name
    _, problems, notes = X.to_csv(saved(literal, entries, merged=("V16:X16",)), literal)
    assert problems == []
    assert notes == ["merged cells V16:X16: only the first cell of a merge holds text"]
    # one over two column names takes the second away
    _, problems, _ = X.to_csv(saved(literal, entries, merged=("L16:M16",)), literal)
    assert problems == ["the header row of the table is not the blank's: M holds '' for 'start'"]


def test_a_changed_shown_cell_is_named_and_written_as_it_is(literal: Path) -> None:
    entries = {
        ("I000000002", "drug"): typed("Sulfasalazine Tablets "),
        ("I000000003", "initial_posting_date"): number("12"),
        ("I000000004", "date_of_update"): day_cell("2021-07-22"),
        ("I000000003", "drug"): typed("=cmd|' /C calc'!A0"),
        ("I000000001", "reason_for_shortage"): ("empty", "", "@"),
        ("I000000001", "statement_type"): typed("recovery"),
    }
    text, problems, notes = X.to_csv(saved(literal, entries), literal)
    assert problems == []
    assert notes == [
        "F19: stored as a number (shown as General), written as the text '12'",
        "H20: stored as a date (shown as mm-dd-yy), written as the text '2021-07-22'",
        "shown cell changed: item I000000001, column reason_for_shortage (K17): '', "
        "the blank holds 'Demand increase for the drug'",
        "shown cell changed: item I000000002, column drug (B18): 'Sulfasalazine Tablets ', "
        "the blank holds 'Sulfasalazine Tablets'",
        "shown cell changed: item I000000003, column drug (B19): \"=cmd|' /C calc'!A0\", "
        "the blank holds \"'=cmd|' /C calc'!A0\"",
        "shown cell changed: item I000000003, column initial_posting_date (F19): '12', "
        "the blank holds '0012'",
    ]
    table = parsed(text)
    assert table["I000000002"]["drug"] == "Sulfasalazine Tablets "  # not put back from the blank
    assert table["I000000003"]["initial_posting_date"] == "12"
    assert table["I000000004"]["date_of_update"] == "2021-07-22"  # the same text: not a change


def test_a_change_far_into_a_long_shown_cell_is_shown(literal: Path) -> None:
    changed = LONG[:20000] + "A" + LONG[20001:]
    entries = {("I000000004", "related_information"): typed(changed)}
    text, problems, notes = X.to_csv(saved(literal, entries), literal)
    assert problems == [] and parsed(text)["I000000004"]["related_information"] == changed
    assert notes == [
        "shown cell changed: item I000000004, column related_information (J20): "
        f"...{changed[19980:20040]!r}..., the blank holds ...{LONG[19980:20040]!r}..."
    ]
    assert X.difference("abc", "abd") == "'abc', the blank holds 'abd'"
    assert X.difference("abc", "abcd") == "'abc', the blank holds 'abcd'"
    assert X.clip("x" * 61) == repr("x" * 60) + "..." and X.clip("x" * 60) == repr("x" * 60)
    assert X.clip("two\nlines") == "'two\\nlines'"  # a message stays on one line


def test_a_change_in_an_entered_cell_is_no_shown_cell_change(literal: Path) -> None:
    text, problems, notes = X.to_csv(saved(literal, ANSWER), literal)
    assert problems == [] and notes == []
    assert text != literal_text() and text.replace('"recovery"', '""') == literal_text()


# --------------------------------------------------------------------------------------------
# The # lines
# --------------------------------------------------------------------------------------------


def with_lines(blank: Path, lines: dict[str, Cell]) -> tuple[str, list[str], list[str]]:
    entries = {(key, "#"): cell for key, cell in lines.items()}
    return X.to_csv(saved(blank, entries), blank)


def test_the_times_are_written_as_typed(literal: Path, outcome: Path) -> None:
    lines = {
        "sitting_start": typed("# sitting_start: 09:30"),
        "sitting_end": typed("# sitting_end: 09:47"),
    }
    text, problems, notes = with_lines(literal, lines)
    assert problems == [] and notes == []
    assert text == literal_text().replace("# sitting_start:", "# sitting_start: 09:30").replace(
        "# sitting_end:", "# sitting_end: 09:47"
    )
    lines = {
        "session_start": typed("# session_start: 13:00; 16:30"),
        "session_end": typed("# session_end: 13:40; 16:55"),
    }
    text, problems, notes = with_lines(outcome, lines)
    assert problems == [] and notes == []
    assert text == outcome_text().replace(
        "# session_start: ", "# session_start: 13:00; 16:30"
    ).replace("# session_end: ", "# session_end: 13:40; 16:55")


@pytest.mark.parametrize(
    ("cell", "line"),
    [
        ("#sitting_start : 09:30", "sitting_start : 09:30"),
        ("#  sitting_start:09:30  ", "sitting_start:09:30  "),
        ("# sitting_start: 9:30 am", "sitting_start: 9:30 am"),
        ("# sitting_start: half past nine", "sitting_start: half past nine"),
    ],
)
def test_a_time_line_typed_with_other_spaces_is_kept(literal: Path, cell: str, line: str) -> None:
    text, problems, notes = with_lines(literal, {"sitting_start": typed(cell)})
    assert problems == [] and notes == []  # whether it is a time is the validator's to say
    assert comments(text)[-2] == line and f'"{cell}"\n' in text


def test_a_time_line_under_another_key_is_a_missing_line(literal: Path) -> None:
    text, problems, notes = with_lines(literal, {"sitting_start": typed("# Sitting start: 09:30")})
    assert (text, problems) == ("", ["# lines missing: # sitting_start"])
    assert notes == ["row 14: a # line that the blank does not hold: '# Sitting start: 09:30'"]


@pytest.mark.parametrize(
    ("cell", "line"),
    [
        ("# sitting_start\uff1a09:30", "sitting_start:09:30"),
        ("# sitting_start: 09\uff1a30", "sitting_start: 09:30"),
        ("# sitting_start: 09:30\uff1b 13:00", "sitting_start: 09:30; 13:00"),
        ("# sitting_start: 09:30\n", "sitting_start: 09:30"),
        ("# sitting_start:\n09:30", "sitting_start: 09:30"),
        ("# sitting_start:_x000D_\n09:30", "sitting_start: 09:30"),
        ("  # sitting_start: 09:30", "sitting_start: 09:30"),
    ],
    ids=[
        "colon after the key",
        "colon in the time",
        "semicolon",
        "break after",
        "break",
        "windows",
        "indent",
    ],
)
@pytest.mark.parametrize("which", ["literal", "outcome"])
def test_a_time_line_is_made_one_plain_line_and_named(
    which: str, cell: str, line: str, request: pytest.FixtureRequest
) -> None:
    blank = request.getfixturevalue(which)
    key = "sitting_start" if which == "literal" else "session_start"
    cell, line = cell.replace("sitting_start", key), line.replace("sitting_start", key)
    text, problems, notes = with_lines(blank, {key: typed(cell)})
    assert problems == []
    row = at(blank, key, "#")[0]
    assert notes[-1] == f"row {row}: the # line is written as {'# ' + line!r}"
    assert comments(text)[-2] == line
    assert "\uff1a" not in text and "\uff1b" not in text
    written_line = text.split("\n")[row - 1]
    assert written_line == (f'"# {line}"' if which == "literal" else f"# {line}")


def test_a_time_typed_beside_its_line_joins_the_line(literal: Path, outcome: Path) -> None:
    for blank, key in ((literal, "sitting_start"), (outcome, "session_start")):
        row = at(blank, key, "#")[0]
        cells = grid(blank)
        cells[row, 2] = number("0.3958333333", "h:mm")  # typed 09:30 into the cell beside the line
        cells[row + 1, 3] = typed(" 09:47 ")
        book = Stored()
        book.add("sheet1", cells)
        text, problems, notes = X.to_csv(book.save(blank.parent / "beside.xlsx"), blank)
        assert problems == []
        assert notes == [
            f"B{row}: stored as a time (shown as h:mm), written as the text '09:30'",
            f"row {row}: the # line is written as '# {key}: 09:30'",
            f"row {row + 1}: the # line is written as '# {key.replace('start', 'end')}: 09:47'",
        ]
        assert comments(text)[-2:] == [f"{key}: 09:30", f"{key.replace('start', 'end')}: 09:47"]


def test_further_sitting_lines_are_kept_in_their_place(literal: Path) -> None:
    def more(rows: list[list[Cell | None]], h: int) -> list[list[Cell | None]]:
        pairs = [[typed("# sitting_start: 09:00")], [typed("# sitting_end: 09:20")]]
        pairs += [[typed("# sitting_start: 13:00")], [typed("# sitting_end: 13:10")]]
        return rows[: h - 2] + pairs + rows[h:]

    text, problems, notes = moved(literal, more)
    assert problems == [] and notes == []
    assert comments(text)[-4:] == [
        "sitting_start: 09:00",
        "sitting_end: 09:20",
        "sitting_start: 13:00",
        "sitting_end: 13:10",
    ]
    assert parsed(text)["I000000001"]["statement_type"] == "recovery"


def test_a_missing_line_stops_the_conversion(literal: Path, outcome: Path) -> None:
    text, problems, _ = moved(literal, lambda rows, h: rows[: h - 1] + rows[h:])
    assert (text, problems) == ("", ["# lines missing: # sitting_end"])
    text, problems, _ = moved(literal, lambda rows, h: rows[:1] + rows[3:])
    assert (text, problems) == ("", ["# lines missing: # annotator; # sheet"])
    # the time typed over the whole cell: the line is gone, and the row is not a # line
    text, problems, _ = with_lines(outcome, {"session_start": typed("13:00")})
    assert text == ""
    assert problems == [
        "row 7, above the table, is not a # line: '13:00'",
        "# lines missing: # session_start",
    ]
    _, problems, _ = moved(literal, lambda rows, h: [[typed("pilot of A1")], *rows])
    assert problems == ["row 1, above the table, is not a # line: 'pilot of A1'"]


def test_a_changed_or_an_added_line_is_named(literal: Path) -> None:
    lines = {
        "annotator": typed("# annotator: A2"),
        "items": typed("# items: 4 "),  # a space more: nothing the validators read
        "seed": typed("# seed:"),
        "values of certainty": typed("# values of certainty: asserted | estimated | sure"),
    }
    text, problems, notes = with_lines(literal, lines)
    assert problems == []
    assert notes == [
        "row 3: # line changed: '# annotator: A2', the blank holds '# annotator: A1'",
        "row 4: # line changed: '# seed:', the blank holds '# seed: 20261001'",
        "row 11: # line changed: '# values of certainty: asserted | estimated | sure', the blank "
        "holds '# values of certainty: asserted | estimated | undetermined | no_statement'",
    ]
    assert comments(text)[2:6] == [
        "annotator: A2",
        "seed:",
        "guide: v1 draft sha256 0000000000000000",
        "items: 4 ",
    ]
    text, problems, notes = moved(
        literal, lambda rows, h: [*rows[:h], [typed("# remark: row 3 is odd")], *rows[h:]]
    )
    assert problems == []
    assert notes == ["row 16: a # line that the blank does not hold: '# remark: row 3 is odd'"]
    assert comments(text)[-1] == "remark: row 3 is odd"


def test_the_keys_of_the_lines_are_read_as_the_validators_read_them() -> None:
    assert X.comment_key("# sitting_start: 09:30") == "sitting_start"
    assert X.comment_key("#sitting_start : 09:30") == "sitting_start"
    assert X.comment_key("  ##  values of hard: 1 or 0") == "values of hard"
    assert X.comment_key("# session_start: 13:00; 16:30") == "session_start"
    assert X.comment_value("# session_start: 13:00; 16:30 ") == "13:00; 16:30"
    assert X.comment_value("# session_start: ") == X.comment_value("# sitting_end:") == ""
    assert X.comment_line(["# a: 1"]) == "# a: 1"
    assert X.comment_line(["# a: ", "", " 1 ", "2"]) == "# a: 1 2"
    assert X.comment_line(["# a: 1 "]) == "# a: 1 "  # nothing to change: the cell as it stands


# --------------------------------------------------------------------------------------------
# The command
# --------------------------------------------------------------------------------------------


def test_the_two_commands(literal: Path, capsys: pytest.CaptureFixture[str]) -> None:
    book, out = literal.parent / "sheet.xlsx", literal.parent / "returned" / "pilot_A1.csv"
    assert X.main(["to-xlsx", str(literal), str(book)]) == 0
    assert capsys.readouterr().out == f"wrote {book}\n"
    filled = saved(literal, {**ANSWER, ("I000000001", "abstain"): number("0")})
    assert X.main(["to-csv", str(filled), "--blank", str(literal), "--out", str(out)]) == 0
    said = capsys.readouterr()
    assert said.out == (
        "note: O17: stored as a number (shown as General), written as the text '0'\n"
        f"wrote {out} (1 notes)\n"
    )
    assert said.err == ""
    row = parsed(out.read_bytes().decode("utf-8"))["I000000001"]
    assert (row["statement_type"], row["abstain"]) == ("recovery", "0")
    assert not out.read_bytes().startswith(b"\xef\xbb\xbf") and b"\r" not in out.read_bytes()


def test_problems_leave_no_file(literal: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = literal.parent / "out.csv"
    filled = saved(literal, {("I000000001", "statement_type"): ("f", "recovery", "General")})
    assert X.main(["to-csv", str(filled), "--blank", str(literal), "--out", str(out)]) == 1
    said = capsys.readouterr()
    assert "error: L17: holds a formula, which is not an entry: '=recovery'" in said.err
    assert f"error: {out} is not written (1 problems)" in said.err
    assert said.out == "" and not out.exists()


def test_an_earlier_return_is_not_replaced_unasked(
    literal: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = literal.parent / "pilot_A1_returned.csv"
    first = saved(literal, ANSWER, name="first.xlsx")
    second = saved(literal, {("I000000001", "statement_type"): typed("none")}, name="second.xlsx")
    arguments = ["--blank", str(literal), "--out", str(out)]
    assert X.main(["to-csv", str(first), *arguments]) == 0
    earlier = out.read_bytes()
    assert X.main(["to-csv", str(first), *arguments]) == 0  # the same sheet again: nothing to do
    assert f"unchanged: {out} already holds this sheet" in capsys.readouterr().out
    assert X.main(["to-csv", str(second), *arguments]) == 2
    assert "exists and holds something else; it is not replaced" in capsys.readouterr().err
    assert out.read_bytes() == earlier
    assert X.main(["to-csv", str(second), *arguments, "--replace"]) == 0
    assert out.read_bytes() != earlier
    assert parsed(out.read_text(encoding="utf-8"))["I000000001"]["statement_type"] == "none"


def test_the_blank_and_the_workbook_are_never_written_over(
    literal: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    filled, before = saved(literal, ANSWER), literal.read_bytes()
    other = literal.parent / "sub" / ".." / literal.name  # the same file by another path
    for out in (literal, other):
        assert (
            X.main(["to-csv", str(filled), "--blank", str(literal), "--out", str(out), "--replace"])
            == 2
        )
        assert "error: --out is the blank sheet" in capsys.readouterr().err
    for name, make in (("link.csv", os.symlink), ("hard.csv", os.link)):
        make(literal, literal.parent / name)
        arguments = ["--blank", str(literal), "--out", str(literal.parent / name), "--replace"]
        assert X.main(["to-csv", str(filled), *arguments]) == 2
        assert "error: --out is the blank sheet" in capsys.readouterr().err
    assert literal.read_bytes() == before
    folder = literal.parent / "returned"
    folder.mkdir()
    assert X.main(["to-csv", str(filled), "--blank", str(literal), "--out", str(folder)]) == 2
    assert "error: --out is a directory" in capsys.readouterr().err
    workbook = filled.read_bytes()
    assert (
        X.main(["to-csv", str(filled), "--blank", str(literal), "--out", str(filled), "--replace"])
        == 2
    )
    assert "error: --out is the filled workbook" in capsys.readouterr().err
    assert filled.read_bytes() == workbook
    assert X.main(["to-xlsx", str(literal), str(literal), "--replace"]) == 2
    assert literal.read_bytes() == before


def test_a_workbook_in_the_way_is_not_replaced_by_a_blank_one(
    literal: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    filled = saved(literal, ANSWER, name="pilot_A1_filled.xlsx")
    before = filled.read_bytes()
    assert X.main(["to-xlsx", str(literal), str(filled)]) == 2
    assert "exists; it is not replaced without --replace" in capsys.readouterr().err
    assert filled.read_bytes() == before
    assert X.main(["to-xlsx", str(literal), str(filled), "--replace"]) == 0
    assert X.to_csv(filled, literal)[0] == literal_text()


def test_a_file_that_is_no_workbook_is_told_so(
    literal: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = literal.parent / "out.csv"
    as_csv = written(literal.parent, "saved_as.csv", literal_text())
    renamed = written(literal.parent, "renamed.xlsx", literal_text())
    other_zip = literal.parent / "numbers.xlsx"
    with zipfile.ZipFile(other_zip, "w") as archive:
        archive.writestr("Index/Document.iwa", "not a workbook")
    for path in (as_csv, renamed, other_zip, literal.parent / "missing.xlsx"):
        assert X.main(["to-csv", str(path), "--blank", str(literal), "--out", str(out)]) == 1
        said = capsys.readouterr().err
        assert said.startswith("error: ") and str(path) in said and "Traceback" not in said
        assert not out.exists()
    with pytest.raises(ValueError, match=r"not an Excel workbook \(.xlsx\)"):
        X.read_xlsx(renamed)


# --------------------------------------------------------------------------------------------
# The validators of the study on converted sheets
# --------------------------------------------------------------------------------------------

LITERAL_ANSWERS = {
    "I000000001": {
        "statement_type": "next_delivery",
        "start": "2021-08-01",
        "end": "2021-08-31",
        "abstain": "0",
        "certainty": "asserted",
        "hard": "0",
    },
    "I000000002": {
        "statement_type": "recovery",
        "start": "2021-10",
        "end": "2021-12",
        "abstain": "0",
        "certainty": "estimated",
        "quote": "Estimated recovery: Q4 2021",
        "hard": "1",
        "note": 'Q4 read as the quarter; "Next Delivery" is the nearer one, see 3.6',
    },
    "I000000003": {
        "statement_type": "none",
        "abstain": "1",
        "abstain_reason": "no_statement",
        "certainty": "no_statement",
    },
    "I000000004": {
        "statement_type": "next_delivery",
        "abstain": "1",
        "abstain_reason": "tbd",
        "certainty": "undetermined",
        "quote": "Supply TBD",
        "distractor_roles": "expiry; past",
        "distractor_quotes": "exp 11/2021 | made 2020",
    },
}
LITERAL_LABELS = {
    "I000000001": {"start": "2021-08-01", "end": "2021-08-31"},
    "I000000002": {"start": "2021-10-01", "end": "2021-12-31"},  # a month: its first, its last day
}
"""The dates the validator reads from the entries above."""
OUTCOME_ANSWERS = {
    "Q0000001": {"verdict_B": "ok", "verdict_A": "ok"},
    "Q0000002": {
        "verdict_B": "error",
        "verdict_A": "ok",
        "codes_B": "O5; O6",
        "true_B_type": "recovered",
        "true_B_lower": "2020-04-02",
        "true_B_upper": "2020-05-01",
        "note_codes": "N1",
        "dou_first_recovered": "2020-04-15",
        "note": "the second capture was skipped, see the trace",
    },
    "Q0000003": {
        "verdict_B": "cannot_tell",
        "verdict_A": "cannot_tell",
        "note": "a gap in the trace",
    },
}


def stored_as(text: str, how: str) -> Cell:
    """An entry as a program stores it: typed into a text cell (``text``), or into a cell that
    is not text (``converted``), where ``1`` is a number, a day a date and a month a date shown
    without its day."""
    if how == "text":
        return typed(text)
    if text in ("0", "1"):
        return number(text)
    if re.fullmatch(r"\d{4}-\d\d-\d\d", text):
        return day_cell(text, "mm-dd-yy" if text < "2021" else "yyyy\\-mm\\-dd")
    if re.fullmatch(r"\d{4}-\d\d", text):
        return day_cell(f"{text}-01", "mmm-yy")
    return typed(text)


SAMPLER = "from analysis.coling import audit_sample, audit_outcomes"


def in_study(python: str, *arguments: str) -> subprocess.CompletedProcess[str]:
    environment = {**os.environ, "PYTHONPATH": str(ROOT)}
    return subprocess.run(
        [python, *arguments], cwd=ROOT, env=environment, capture_output=True, text=True, check=False
    )


def study_python() -> str | None:
    """The interpreter of the study's environment: the one that ``COLING_STUDY_PYTHON`` names,
    or the ``.venv`` of this checkout, or the ``.venv`` of the checkout this one is a worktree
    of. It has to import the sampler."""
    named = os.environ.get("COLING_STUDY_PYTHON")
    places = [Path(named)] if named else []
    places.append(ROOT / ".venv" / "bin" / "python")
    link = ROOT / ".git"
    if link.is_file():  # a worktree: the file names a directory under the main checkout's .git
        directory = Path(link.read_text(encoding="utf-8").partition("gitdir:")[2].strip())
        places.append(directory.parents[2] / ".venv" / "bin" / "python")
    for place in places:
        if place.is_file() and in_study(str(place), "-c", SAMPLER).returncode == 0:
            return str(place)
    return None


@pytest.fixture(scope="module")
def study() -> str:
    found = study_python()
    if found is None:
        pytest.skip("the interpreter of the study's environment is not found")
    return found


READ_LABELS = """
import json, sys
from pathlib import Path
from analysis.coling import audit_sample as S
checked = S.check_file(Path(sys.argv[1]), Path(sys.argv[2]))
print(json.dumps({"labels": {i: l.entered() for i, l in checked.labels.items()},
                  "errors": checked.errors, "minutes": checked.minutes}))
"""
READ_ROWS = """
import json, sys
from pathlib import Path
from analysis.coling import audit_outcomes as AO
meta, rows = AO.read_sheet(Path(sys.argv[1]))
print(json.dumps({"rows": rows, "minutes": AO.session_minutes(meta)}))
"""


@pytest.mark.parametrize("how", ["text", "converted"])
def test_the_literal_validator_reads_the_labels_that_were_typed(
    literal: Path, study: str, how: str
) -> None:
    entries = {
        (item, column): stored_as(text, how)
        for item, row in LITERAL_ANSWERS.items()
        for column, text in row.items()
    }
    entries[("sitting_start", "#")] = typed("# sitting_start: 09:30")
    entries[("sitting_end", "#")] = typed(
        "# sitting_end\uff1a09:47" if how == "converted" else "# sitting_end: 09:47"
    )
    out = literal.parent / "returned.csv"
    assert (
        X.main(["to-csv", str(saved(literal, entries)), "--blank", str(literal), "--out", str(out)])
        == 0
    )
    run = in_study(
        study, "-m", "analysis.coling.audit_sample", "validate", str(out), "--blank", str(literal)
    )
    assert run.returncode == 0, run.stdout + run.stderr
    assert "4 rows, 4 valid, 0 errors" in run.stdout and "sittings: 17 minutes" in run.stdout
    read = json.loads(in_study(study, "-c", READ_LABELS, str(out), str(literal)).stdout)
    assert read["errors"] == [] and read["minutes"] == 17
    for item, row in LITERAL_ANSWERS.items():
        typed_row = {column: "" for column in LITERAL_ENTERED} | row | LITERAL_LABELS.get(item, {})
        typed_row["hard"] = typed_row["hard"] or "0"
        assert read["labels"][item] == typed_row, item


def test_the_literal_validator_adds_up_further_sittings(literal: Path, study: str) -> None:
    entries = {
        (item, column): typed(text)
        for item, row in LITERAL_ANSWERS.items()
        for column, text in row.items()
    }
    rows, at_header = rows_of(literal, entries), len(LITERAL_META)
    lines = ("# sitting_start: 09:30", "# sitting_end: 09:47", "# sitting_start: 13:00")
    pairs = [[typed(line)] for line in lines] + [
        [typed("# sitting_end:"), number("0.5625", "h:mm")]
    ]
    filled = from_rows(literal, [*rows[: at_header - 2], *pairs, *rows[at_header:]])
    out = literal.parent / "returned.csv"
    assert X.main(["to-csv", str(filled), "--blank", str(literal), "--out", str(out)]) == 0
    run = in_study(
        study, "-m", "analysis.coling.audit_sample", "validate", str(out), "--blank", str(literal)
    )
    assert run.returncode == 0, run.stdout + run.stderr
    assert "4 rows, 4 valid, 0 errors" in run.stdout and "sittings: 47 minutes" in run.stdout


@pytest.mark.parametrize("how", ["text", "converted"])
def test_the_outcome_validator_reads_the_verdicts_that_were_typed(
    outcome: Path, study: str, how: str
) -> None:
    entries = {
        (item, column): stored_as(text, how)
        for item, row in OUTCOME_ANSWERS.items()
        for column, text in row.items()
    }
    entries[("session_start", "#")] = typed("# session_start: 13:00; 16:30")
    entries[("session_end", "#")] = typed(
        "# session_end: 13:40\uff1b 16:55" if how == "converted" else "# session_end: 13:40; 16:55"
    )
    out = outcome.parent / "returned.csv"
    assert (
        X.main(["to-csv", str(saved(outcome, entries)), "--blank", str(outcome), "--out", str(out)])
        == 0
    )
    run = in_study(
        study, "-m", "analysis.coling.audit_outcomes", "validate", str(out), "--blank", str(outcome)
    )
    assert run.returncode == 0, run.stdout + run.stderr
    assert "3 rows, 0 problems, 0 warnings" in run.stdout
    read = json.loads(in_study(study, "-c", READ_ROWS, str(out)).stdout)
    assert read["minutes"] == 65
    blank_rows = parsed(outcome_text())
    for row in read["rows"]:
        typed_row = {column: "" for column in OUTCOME_ENTERED} | OUTCOME_ANSWERS[row["audit_id"]]
        assert {column: row[column] for column in OUTCOME_ENTERED} == typed_row
        assert all(row[c] == blank_rows[row["audit_id"]][c].strip() for c in OUTCOME_SHOWN)


def test_what_the_validators_reject_is_still_there_to_reject(literal: Path, study: str) -> None:
    # the converter guesses no label: TRUE stays TRUE, and the validator says so
    entries = {
        (item, column): stored_as(text, "text")
        for item, row in LITERAL_ANSWERS.items()
        for column, text in row.items()
    }
    entries[("I000000001", "hard")] = ("b", "0", "General")
    entries[("I000000003", "abstain")] = ("b", "1", "General")
    out = literal.parent / "returned.csv"
    assert (
        X.main(["to-csv", str(saved(literal, entries)), "--blank", str(literal), "--out", str(out)])
        == 0
    )
    run = in_study(
        study, "-m", "analysis.coling.audit_sample", "validate", str(out), "--blank", str(literal)
    )
    assert run.returncode == 1
    assert "abstain 'TRUE' is not one of 0, 1" in run.stdout
    assert "hard 'FALSE' is not one of 0, 1" in run.stdout


# --------------------------------------------------------------------------------------------
# The workbooks that were handed out (read only)
# --------------------------------------------------------------------------------------------


@pytest.mark.parametrize(("name", "blank"), HANDED_OUT, ids=[name for name, _ in HANDED_OUT])
def test_a_handed_out_workbook(name: str, blank: Path) -> None:
    path = HANDOUT / name
    if not path.is_file() or not blank.is_file():
        pytest.skip("the handed-out workbook or its blank sheet is not on disk")
    before = path.read_bytes()
    text, problems, notes = X.to_csv(path, blank)
    assert text.encode("utf-8") == blank.read_bytes() and problems == [] and notes == []
    formats, _ = text_cells(path)
    assert set(formats.values()) == {"49"}
    assert all(ref in formats for ref in cells_to_fill(blank))
    base = X.read_csv_sheet(blank)
    wanted = {
        X.column_letter(base.header.index(column) + 1): '"' + ",".join(values) + '"'
        for column, values in X.value_lists(base).items()
    }
    lists = lists_of(path)
    assert {letter: rule["values"] for letter, rule in lists.items()} == wanted
    first, last = base.header_at + 2, len(base.rows)
    for letter, rule in lists.items():
        assert rule["sqref"] == f"{letter}{first}:{letter}{last}"
        assert (rule["type"], rule["allowBlank"], rule["showErrorMessage"]) == ("list", "1", "0")
    with zipfile.ZipFile(path) as archive:
        core = archive.read("docProps/core.xml").decode("utf-8")
    people = dict(re.findall(r"<(?:dc|cp):(creator|lastModifiedBy)[^>]*>([^<]*)</", core))
    # an element written empty or left out names nobody (the library does either)
    assert set(people) <= {"creator", "lastModifiedBy"} and set(people.values()) <= {""}
    assert path.read_bytes() == before


# --- the sheet of the minimal-pair audit (task C) ---

PAIR_MARKS = (
    "ok_type",
    "ok_interval",
    "ok_certainty",
    "ok_stale",
    "ok_distractors",
    "minimal",
    "attested",
    "natural",
)
PAIR_ENTERED = (*PAIR_MARKS, "correct_value", "note")


def pairs_text() -> str:
    """A blank sheet of the minimal-pair audit: entries of several lines, ten columns to fill."""
    shown = study_constants("minimal_pairs.py", {"SHOWN"})["SHOWN"]
    out = io.StringIO(newline="")
    writer = csv.writer(out, lineterminator="\n", quoting=csv.QUOTE_ALL)
    for line in (
        "# task: C minimal-pair audit (AUDIT_GUIDE.md section 7)",
        "# sheet: pairs",
        "# annotator: A1",
        "# values of minimal: 1 when the edit changes only the named factor, else 0",
        "# values of note: free text; required when minimal or attested is 0",
        "# sitting_start:",
        "# sitting_end:",
    ):
        writer.writerow([line])
    writer.writerow([*shown, *PAIR_ENTERED])
    for n in ("1", "2"):
        entry = f'Entry\n- Drug: Name{n} Injection\n- Related information: "Next release TBD"'
        cells = dict.fromkeys(shown, "")
        cells |= {"pair_id": f"P{n}", "seed_id": f"S{n}", "factor": "certainty", "level": "tbd"}
        cells |= {"seed_entry": entry, "edited_entry": entry, "attested_text": "Next release TBD"}
        writer.writerow([*cells.values(), *[""] * len(PAIR_ENTERED)])
    return out.getvalue()


def test_a_minimal_pair_sheet_has_its_columns_to_fill_and_its_lists(tmp_path: Path) -> None:
    checks = study_constants("minimal_pairs.py", {"SHOWN", "CHECKS"})
    assert (*checks["CHECKS"], "minimal", "attested", "natural") == PAIR_MARKS
    blank = written(tmp_path, "pairs_A1.csv", pairs_text())
    sheet = X.read_csv_sheet(blank)
    assert X.first_entered(sheet.header) == len(checks["SHOWN"])
    assert X.value_lists(sheet) == dict.fromkeys(PAIR_MARKS, ("1", "0"))
    book = blank.with_suffix(".xlsx")
    X.to_xlsx(blank, book)
    letters = [X.column_letter(sheet.header.index(name) + 1) for name in PAIR_MARKS]
    assert sorted(lists_of(book)) == sorted(letters)
    assert len(cells_to_fill(blank)) == 2 * len(PAIR_ENTERED) + 2
    back, problems, notes = X.to_csv(book, blank)
    assert back.encode("utf-8") == blank.read_bytes() and not problems and not notes
