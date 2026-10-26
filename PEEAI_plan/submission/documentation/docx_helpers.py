"""Formatting helpers for build_doc.py (python-docx 1.1.2): page setup, text, lists, tables, figures.

Command lines and identifiers use two monospace styles, the paragraph style "Code" and the
character style "Inline Code". The writing-rule check (../check_writing.py) skips text in those
two styles and nothing else.
"""

import datetime
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.run import Run

PRIMARY = RGBColor(0x08, 0x45, 0x7E)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
PLACEHOLDER = RGBColor(0xB3, 0x47, 0x00)
LINK = RGBColor(0x05, 0x63, 0xC1)
TEXT_WIDTH_IN = 6.5


def new_document(title: str, author: str, subject: str, keywords: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))
    props = doc.core_properties
    props.title, props.author, props.last_modified_by = title, author, author
    props.subject, props.keywords, props.comments = subject, keywords, ""
    props.created = props.modified = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0)
    props.revision = 1
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for level, size in ((1, 16), (2, 13), (3, 11.5)):
        heading = doc.styles[f"Heading {level}"]
        heading.font.name = "Calibri"
        heading.font.size = Pt(size)
        heading.font.color.rgb = PRIMARY
    code = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    code.base_style = normal
    code.font.name = "Consolas"
    code.font.size = Pt(9)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.left_indent = Inches(0.25)
    inline = doc.styles.add_style("Inline Code", WD_STYLE_TYPE.CHARACTER)
    inline.font.name = "Consolas"
    inline.font.size = Pt(9.5)
    add_page_number_footer(section)
    return doc


def add_page_number_footer(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def rich(paragraph, text: str) -> None:
    """Add text where **bold** spans are bold, `code` spans are inline code, <<url>> spans are
    clickable links, [[...]] spans are marked result placeholders and {{...}} spans are short
    placeholder tokens for table cells."""
    tokens = []
    rest = text
    while rest:
        starts = [(rest.find(m), m) for m in ("**", "[[", "{{", "`", "<<") if rest.find(m) >= 0]
        if not starts:
            tokens.append(("plain", rest))
            break
        index, mark = min(starts)
        if index:
            tokens.append(("plain", rest[:index]))
        close = {"**": "**", "[[": "]]", "{{": "}}", "`": "`", "<<": ">>"}[mark]
        kind = {"**": "bold", "[[": "placeholder", "{{": "token", "`": "code", "<<": "link"}[mark]
        end = rest.find(close, index + len(mark))
        tokens.append((kind, rest[index + len(mark) : end]))
        rest = rest[end + len(close) :]
    for kind, piece in tokens:
        if kind in ("placeholder", "token"):
            text = f"[Wave 3 result: {piece} pending]" if kind == "placeholder" else f"[{piece}]"
            run = paragraph.add_run(text)
            run.bold = True
            run.font.color.rgb = PLACEHOLDER
        elif kind == "code":
            paragraph.add_run(piece, style="Inline Code")
        elif kind == "link":
            hyperlink(paragraph, piece)
        else:
            run = paragraph.add_run(piece)
            run.bold = kind == "bold"


def hyperlink(paragraph, url: str) -> None:
    """Add url as a clickable external link, blue and underlined, showing the url itself."""
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    element = OxmlElement("w:r")
    link.append(element)
    paragraph._p.append(link)
    run = Run(element, paragraph)
    run.text = url
    run.font.color.rgb = LINK
    run.font.underline = True


def para(doc: Document, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    rich(paragraph, text)
    return paragraph


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        para(doc, item, "List Bullet")


def code_block(doc: Document, lines: list[str]) -> None:
    """Command lines, one paragraph each, in the monospace "Code" style."""
    for line in lines:
        doc.add_paragraph(line, style="Code")
    doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def restart_numbering(doc: Document, first_paragraph) -> None:
    """Start a new List Number sequence at first_paragraph (a fresh w:num on the list's abstract)."""
    numbering = doc.part.numbering_part.element
    style_num = doc.styles["List Number"].element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(style_num).abstractNumId.val
    num = numbering.add_num(abstract)
    override = num.add_lvlOverride(ilvl=0)
    override.add_startOverride(1)
    pPr = first_paragraph._p.get_or_add_pPr()
    numPr = pPr.get_or_add_numPr()
    numPr.get_or_add_ilvl().val = 0
    numPr.get_or_add_numId().val = num.numId


def numbered(doc: Document, items: list[str]) -> None:
    first = None
    for item in items:
        paragraph = para(doc, item, "List Number")
        if first is None:
            first = paragraph
            restart_numbering(doc, paragraph)
        else:
            pPr = paragraph._p.get_or_add_pPr()
            numPr = pPr.get_or_add_numPr()
            numPr.get_or_add_ilvl().val = 0
            numPr.get_or_add_numId().val = first._p.pPr.numPr.numId.val


def _shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    tcPr.append(shading)


def table(doc: Document, header: list[str], rows: list[list[str]], widths_in: list[float]) -> None:
    assert len(header) == len(widths_in) and abs(sum(widths_in) - TEXT_WIDTH_IN) < 0.05
    grid = doc.add_table(rows=1, cols=len(header))
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    for row_cells, values, is_header in [(grid.rows[0].cells, header, True)] + [
        (grid.add_row().cells, row, False) for row in rows
    ]:
        for cell, value, width in zip(row_cells, values, widths_in):
            cell.width = Inches(width)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            rich(paragraph, f"**{value}**" if is_header else value)
            for item in paragraph.iter_inner_content():
                for run in getattr(item, "runs", [item]):
                    run.font.size = Pt(9.5)
            if is_header:
                _shade(cell, "E6EEF6")
    for column, width in zip(grid.columns, widths_in):
        column.width = Inches(width)
    _keep_rows_whole(grid)
    doc.add_paragraph()


def _keep_rows_whole(grid) -> None:
    """Word repeats the header row on each page a table spans, and never splits a row."""
    for index, row in enumerate(grid.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if index == 0:
            trPr.append(OxmlElement("w:tblHeader"))


def figure(doc: Document, image: Path, caption: str, width_in: float = TEXT_WIDTH_IN) -> None:
    doc.add_picture(str(image), width=Inches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True  # the caption stays with its picture
    paragraph = doc.add_paragraph(style="Caption")
    rich(paragraph, caption)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def neutral_app_properties(path: Path) -> None:
    """Replace python-docx's default docProps/app.xml, which names Microsoft Word and its
    Normal.dotm template, with an empty extended-properties part."""
    minimal = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"/>'
    ).encode("utf-8")
    temporary = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = minimal if item.filename == "docProps/app.xml" else source.read(item.filename)
            target.writestr(item, data)
    temporary.replace(path)
